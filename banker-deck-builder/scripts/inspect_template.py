"""Inspect a deck/template before building: slides, layouts, placeholders, pictures, theme fonts and colors.

Usage: python3 inspect_template.py deck.pptx [--content-slide N]

Prints what you need to write theme.json for deck_kit:
  - every slide with its layout and the text in its title-like placeholders
  - for the chosen content slide (default: first slide whose layout has >= 3 placeholders): each placeholder's
    idx, type, name and position, so you can set title_idx / takeaway_idx / footer_idx / drop_idx
  - pictures on every slide (names you can pass to deck_kit.copy_picture, e.g. logos already in the deck)
  - theme fonts and the colors actually used in shapes on the content slide (candidates for the palette)
"""
import re
import sys
import zipfile
from collections import Counter

from pptx import Presentation
from pptx.util import Emu


def inches(v):
    return round(Emu(v).inches, 2) if v is not None else None


def main(path, content_idx=None):
    prs = Presentation(path)
    print(f"Slide size: {inches(prs.slide_width)} x {inches(prs.slide_height)} in; {len(prs.slides)} slides\n")
    print("== Slides")
    for i, s in enumerate(prs.slides):
        phs = {ph.placeholder_format.idx: (ph.text_frame.text[:70] if ph.has_text_frame else "")
               for ph in s.placeholders}
        title = next((t for t in phs.values() if t), "")
        print(f"  [{i}] layout='{s.slide_layout.name}' shapes={len(s.shapes)} first-text='{title}'")
    if content_idx is None:
        content_idx = next((i for i, s in enumerate(prs.slides) if len(s.slide_layout.placeholders) >= 3), 0)
    s = prs.slides[content_idx]
    print(f"\n== Placeholders on slide [{content_idx}] (layout '{s.slide_layout.name}')")
    for ph in s.placeholders:
        pf = ph.placeholder_format
        txt = ph.text_frame.text[:60].replace("\n", " / ") if ph.has_text_frame else ""
        print(f"  idx={pf.idx:<3} type={str(pf.type):<22} name='{ph.name}' pos=({inches(ph.left)},{inches(ph.top)}) "
              f"size=({inches(ph.width)}x{inches(ph.height)}) text='{txt}'")
    print("\n== Layout placeholders (what a new slide on this layout will get)")
    for ph in s.slide_layout.placeholders:
        pf = ph.placeholder_format
        print(f"  idx={pf.idx:<3} type={str(pf.type):<22} name='{ph.name}' pos=({inches(ph.left)},{inches(ph.top)})")
    print("\n== Pictures (usable with deck_kit.copy_picture)")
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.shape_type == 13 or sh._element.tag.endswith("}pic"):
                print(f"  slide [{i}] '{sh.name}' at ({inches(sh.left)},{inches(sh.top)}) {inches(sh.width)}x{inches(sh.height)}")
    with zipfile.ZipFile(path) as z:
        themes = [n for n in z.namelist() if re.match(r"ppt/theme/theme\d+\.xml", n)]
        if themes:
            t = z.read(themes[0]).decode("utf8", "ignore")
            fonts = re.findall(r'<a:(major|minor)Font><a:latin typeface="([^"]*)"', t)
            print(f"\n== Theme fonts: {fonts}  (Office defaults like Aptos/Calibri mean inserted charts will not inherit "
                  f"the deck font unless set explicitly)")
        sx = z.read(f"ppt/slides/slide{content_idx + 1}.xml").decode("utf8", "ignore") \
            if f"ppt/slides/slide{content_idx + 1}.xml" in z.namelist() else ""
        cols = Counter(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', sx))
        fonts_used = Counter(re.findall(r'typeface="([^"]+)"', sx))
        print(f"== Colors used on slide [{content_idx}] (most common first): {cols.most_common(12)}")
        print(f"== Fonts used on slide [{content_idx}]: {fonts_used.most_common(5)}")
    print("\nNext: copy assets/theme.example.json to the work folder as theme.json and set template_slide, the "
          "placeholder idx values, palette roles and font from the output above.")


if __name__ == "__main__":
    a = sys.argv[1:]
    ci = int(a[a.index("--content-slide") + 1]) if "--content-slide" in a else None
    main(a[0], ci)
