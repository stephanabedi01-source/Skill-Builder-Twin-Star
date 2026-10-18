"""Render Excel model ranges as PNG images for appendix 'model output' pages (they look like the real model,
in the model's own formatting, which reads far less AI-made than rebuilt tables).

Usage: python3 render_xl.py jobs.json [job ...]

jobs.json:
{
  "workbook": "model/Model_v16_recalculated.xlsx",   # MUST hold cached values (recalculated by LibreOffice/Excel)
  "outdir": "xl_img",
  "dpi": 250,
  "jobs": {
    "dcf_top":  {"sheet": "DCF", "range": "C5:N17"},
    "dcf_grid": {"sheet": "DCF", "range": "C68:K76", "overrides": {"G72": 2602.2}}
  }
}

How it works: loads the workbook with data_only=True (values, keeps styles), deletes every other sheet, sets the
print area to the range (landscape, fit to one page, no gridlines/headings/margins), converts with LibreOffice
and trims the PDF page to a tight PNG.

Gotchas:
- openpyxl-written files have NO cached values until recalculated. Render from a recalculated copy.
- Excel data tables (What-If grids) are not evaluated by LibreOffice; their anchor cell keeps a stale value.
  Compute those cells and pass them in "overrides".
- Never render a range that shows a number contradicting the main slides (e.g. an implied EV built on a different
  EBITDA). Crop the range or fix the model first.
"""
import json
import os
import shutil
import subprocess
import sys

from openpyxl import load_workbook
from openpyxl.worksheet.page import PageMargins


def render(cfg, name):
    job = cfg["jobs"][name]
    wb = load_workbook(cfg["workbook"], data_only=True)
    for n in wb.sheetnames:
        if n != job["sheet"]:
            del wb[n]
    ws = wb[job["sheet"]]
    for addr, v in job.get("overrides", {}).items():
        ws[addr].value = v
    ws.print_area = job["range"]
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False
    ws.print_options.headings = False
    ws.page_margins = PageMargins(left=0.1, right=0.1, top=0.1, bottom=0.1, header=0, footer=0)
    for hf in (ws.oddHeader, ws.oddFooter):
        hf.left.text = hf.center.text = hf.right.text = None
    out = cfg.get("outdir", "xl_img")
    tmp = os.path.join(out, f"_{name}")
    os.makedirs(tmp, exist_ok=True)
    x = os.path.join(tmp, f"{name}.xlsx")
    wb.save(x)
    subprocess.run(["timeout", "120", "soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, x],
                   capture_output=True)
    base = os.path.join(tmp, "p")
    subprocess.run(["pdftoppm", "-r", str(cfg.get("dpi", 250)), "-png", "-f", "1", "-l", "1", "-singlefile",
                    os.path.join(tmp, f"{name}.pdf"), base], check=True)
    dst = os.path.join(out, f"{name}.png")
    subprocess.run(["convert", base + ".png", "-trim", "+repage", "-bordercolor", "white", "-border", "6", dst],
                   check=True)
    shutil.rmtree(tmp, ignore_errors=True)
    print("rendered", name, dst)


if __name__ == "__main__":
    cfg = json.load(open(sys.argv[1]))
    os.makedirs(cfg.get("outdir", "xl_img"), exist_ok=True)
    for j in (sys.argv[2:] or list(cfg["jobs"])):
        render(cfg, j)
