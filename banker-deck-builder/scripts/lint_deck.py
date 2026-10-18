"""Lint a deck for the writing and hygiene rules the user cares about. Run before every delivery.

Usage: python3 lint_deck.py deck.pptx [--takeaway-idx 16] [--max-takeaway 120] [--allow-pronouns]

Flags (slide number, shape name, text):
  - personal pronouns on slides (we / our / us / you / your / I / my): slides speak in the third person
  - em dashes and en dashes used as punctuation (obvious AI tell); numeric ranges should read "4.9x to 6.4x"
  - stock AI phrasing (delve, seamless, robust, unlock, landscape, game-changer, ...)
  - leftover placeholders: TBD, XX, [source], [Your Name], lorem, "Click to add"
  - negative numbers written $(152)M instead of ($152M)
  - takeaway lines longer than --max-takeaway characters (they must fit on one line)
  - double spaces
Exit code 1 if anything is flagged, so it can gate a delivery.
"""
import re
import sys

from pptx import Presentation

PRONOUNS = re.compile(r"\b(we|our|ours|us|you|your|yours)\b", re.I)
PRONOUN_CAPS = re.compile(r"(?<![\w\-/.&])(I|my|My)(?![\w\-/.&])")  # skips names like "O-I"
DASHES = re.compile(r"[—–]")
AI_WORDS = re.compile(r"\b(delve|seamless(ly)?|robust|unlock(s|ing)?|game[- ]chang\w*|landscape|tapestry|"
                      r"navigate the|in today's|it's worth noting|testament to|leverag(e|ing) synergies)\b", re.I)
PLACEHOLDERS = re.compile(r"(\bTBD\b|\bXX+\b|\[source\]|\[your name\]|lorem|click to add|placeholder)", re.I)
BAD_NEG = re.compile(r"\$\(\d")


def texts(shape):
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            yield "".join(r.text for r in p.runs)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            yield from texts(s)


def main():
    a = sys.argv[1:]
    deck = a[0]
    tk_idx = int(a[a.index("--takeaway-idx") + 1]) if "--takeaway-idx" in a else 16
    max_tk = int(a[a.index("--max-takeaway") + 1]) if "--max-takeaway" in a else 120
    allow_pron = "--allow-pronouns" in a
    prs = Presentation(deck)
    hits = []
    for n, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            for t in texts(sh):
                if not t.strip():
                    continue
                checks = [(DASHES, "dash"), (AI_WORDS, "ai-phrasing"), (PLACEHOLDERS, "placeholder"),
                          (BAD_NEG, "negative-format")]
                if not allow_pron:
                    checks += [(PRONOUNS, "pronoun"), (PRONOUN_CAPS, "pronoun")]
                for rx, tag in checks:
                    if tag == "dash" and t.strip() in ("\u2013", "\u2014", "(\u2013)"):
                        continue  # banker convention: a lone dash marks a nil cell
                    m = rx.search(t)
                    if m:
                        hits.append((n, tag, sh.name, t[:110]))
                if "  " in t.strip():
                    hits.append((n, "double-space", sh.name, t[:110]))
        for ph in s.placeholders:
            if ph.placeholder_format.idx == tk_idx and ph.has_text_frame and len(ph.text_frame.text) > max_tk:
                hits.append((n, f"takeaway>{max_tk}", ph.name, f"{len(ph.text_frame.text)} chars"))
    for h in hits:
        print(f"slide {h[0]:>2} [{h[1]}] {h[2]}: {h[3]}")
    print(f"{len(hits)} issue(s)")
    sys.exit(1 if hits else 0)


if __name__ == "__main__":
    main()
