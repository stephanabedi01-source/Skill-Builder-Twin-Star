"""deck_kit: building blocks for banker-style slides on an existing PowerPoint template (python-pptx).

Import it from a per-slide builder script:

    import sys; sys.path.insert(0, "<skill>/scripts")
    from deck_kit import *
    load_theme("theme.json")                       # palette, font, template geometry for THIS deck
    prs = Presentation("deck_in.pptx")
    s = new_content_slide(prs, "Valuation Summary", "One-line takeaway that states the claim",
                          notes="Notes: (1) ...", sources="Sources: ...")
    ...
    move_slide(prs, s, 6)
    prs.save("deck_out.pptx")

Everything visual reads the THEME, so the same builder works for a different palette, font or master.
Units are inches unless stated. Colors are 6-digit hex strings without '#'.
"""
import copy
import json

from lxml import etree
from pptx import Presentation  # noqa: F401  (re-exported for builders)
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

__all__ = [n for n in dir() if not n.startswith("_")]  # re-export pptx enums for builders

# ===================================================================== theme
# Role-based tokens. Values are the Project Print deck's (evidence, not constants): override per deck.
THEME = {
    "font": "Garamond",
    "ink": "1A1A24",          # body text, dark bars, rules, totals
    "accent": "D92332",       # emphasis text, lines, pill borders, markers (the subject company / distress only)
    "accent_fill": "C00000",  # accent bar fills
    "grey_txt": "595959",     # captions, secondary labels, footnotes
    "grey_mid": "7F8590",     # secondary series, axes, connectors
    "grey_lt": "BFC3C9",      # peer / reference bars
    "hair": "D1D1D1",         # hairline table rules
    "band": "E9E9E9",         # shaded ranges / bands behind charts, highlighted table rows
    "callout_fill": "D6D6D6", # commentary-callout fill (text callouts beside evidence)
    "white": "FFFFFF",
    # template geometry: which existing slide supplies the layout + footer, and the placeholder idx values
    "template_slide": 2,      # 0-based index of an existing CONTENT slide to clone layout/footer from
    "title_idx": 21,
    "takeaway_idx": 16,
    "drop_idx": [20],         # placeholders on the layout that must be removed on new slides
    "footer_idx": 18,         # footer placeholder that carries Notes / Sources
    "footer_top": 6.93,
    "footer_size": 10,
    "copy_ph_types": ["sldNum", "ftr"],
    # page grid (13.33 x 7.5 in): left panel ~60%, right panel ~35%
    "grid": {"L0": 0.23, "R1": 13.11, "LP1": 8.33, "RP0": 8.62, "ROW1": 1.30,
             "closing_y": 6.43, "closing_h": 0.36, "callout_y": 5.62},
}
G = THEME["grid"]


def load_theme(path_or_dict):
    """Merge a theme JSON file (or dict) into THEME. Unknown keys are kept, so builders can add their own."""
    d = path_or_dict if isinstance(path_or_dict, dict) else json.load(open(path_or_dict))
    grid = d.pop("grid", None)
    THEME.update(d)
    if grid:
        THEME["grid"].update(grid)
    return THEME


def c(role):
    """Color for a role name ('ink', 'accent', ...) or pass a raw hex through."""
    return THEME.get(role, role)


def rgb(h):
    return RGBColor.from_string(c(h))


# ===================================================================== text
def style_run(run, size, bold=False, italic=False, color="ink", sup=False):
    f = run.font
    f.name = THEME["font"]
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    if sup:
        run._r.get_or_add_rPr().set("baseline", "30000")


def fill_tf(tf, paras, size, align=PP_ALIGN.LEFT, space_after=0, line_spacing=None):
    """paras = list of paragraphs; each paragraph = list of (text, opts) runs.
    opts keys: b, i, c (role or hex), sup, size."""
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, o in runs:
            r = p.add_run()
            r.text = text
            style_run(r, o.get("size", size), o.get("b", False), o.get("i", False), o.get("c", "ink"),
                      o.get("sup", False))


def textbox(slide, x, y, w, h, paras, size=11, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None,
            space_after=0, wrap=True, fill=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    fill_tf(tf, paras, size, align, space_after)
    if fill:
        tb.fill.solid()
        tb.fill.fore_color.rgb = rgb(fill)
    if name:
        tb.name = name
    return tb


# ===================================================================== shapes
def no_effects(shape):
    """Strip theme effects (shadows/glow) so shapes render flat and clean."""
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    spPr.append(etree.SubElement(spPr, qn("a:effectLst")))


def box(slide, x, y, w, h, fill="white", line=None, dash=None, line_w=0.75, radius=None, name=None,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = dash
    no_effects(s)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.word_wrap = True
    if name:
        s.name = name
    return s


def line(slide, x1, y1, x2, y2, color="ink", width=1.0, dash=None, name=None):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = rgb(color)
    cn.line.width = Pt(width)
    if dash:
        cn.line.dash_style = dash
    if name:
        cn.name = name
    return cn


def arrow(slide, x1, y1, x2, y2, color="ink", width=1.0, name=None):
    cn = line(slide, x1, y1, x2, y2, color, width, name=name)
    t = etree.SubElement(cn.line._get_or_add_ln(), qn("a:tailEnd"))
    t.set("type", "triangle")
    t.set("w", "sm")
    t.set("len", "sm")
    return cn


def section_header(slide, x, y, w, text, name, sup=None, size=14):
    """Bold header with a full-width rule 0.29in below its top. The deck's universal panel header."""
    runs = [(text, {"b": True})] + ([(sup, {"sup": True, "b": True})] if sup else [])
    textbox(slide, x, y, w, 0.24, [runs], size=size, name=name)
    line(slide, x, y + 0.29, x + w, y + 0.29, "ink", 1.0, name=name + " rule")


def caption(slide, x, y, w, text, name, sup=None, size=10):
    """Italic grey chart caption (units / metric) under a section header."""
    runs = [(text, {"i": True, "c": "grey_txt"})] + ([(sup, {"sup": True, "i": True, "c": "grey_txt"})] if sup else [])
    return textbox(slide, x, y, w, 0.2, [runs], size=size, name=name)


def callout(slide, x, y, w, h, runs, name, size=10.5, align=PP_ALIGN.CENTER):
    """Commentary callout: grey fill, accent round-dot border. Holds the 'so what' next to the evidence."""
    b = box(slide, x, y, w, h, fill="callout_fill", line="accent", dash=MSO_LINE_DASH_STYLE.ROUND_DOT,
            line_w=1.0, radius=0.08, name=name)
    b.text_frame.margin_left = b.text_frame.margin_right = Inches(0.08)
    b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_tf(b.text_frame, [runs], size, align=align)
    return b


def pill(slide, x, y, w, h, runs, name, line_color="accent", fill="white", size=10, dotted=True):
    """Small rounded flag/label tied to a point on a chart (white fill keeps small text readable)."""
    p = box(slide, x, y, w, h, fill=fill, line=line_color, dash=MSO_LINE_DASH_STYLE.ROUND_DOT if dotted else None,
            line_w=1.0, radius=0.5, name=name)
    p.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p.text_frame.margin_left = p.text_frame.margin_right = Inches(0.03)
    p.text_frame.margin_top = p.text_frame.margin_bottom = 0
    fill_tf(p.text_frame, [runs], size, align=PP_ALIGN.CENTER)
    return p


def closing_line(slide, runs, name="Slide conclusion", size=11.5):
    """Full-width dashed ink box at the bottom: the one-sentence conclusion that hands off to the next slide."""
    g = THEME["grid"]
    b = box(slide, g["L0"], g["closing_y"], g["R1"] - g["L0"], g["closing_h"], fill="white", line="ink",
            dash=MSO_LINE_DASH_STYLE.DASH, line_w=0.75, radius=0.2, name=name)
    b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_tf(b.text_frame, [runs], size, align=PP_ALIGN.CENTER)
    return b


def brace(slide, x, y_top, y_bot, name, width=0.09, color="ink"):
    """Vertical right-brace that ties a gap between two values to its label (e.g., 'runoff $663M')."""
    br = slide.shapes.add_shape(MSO_SHAPE.RIGHT_BRACE, Inches(x), Inches(y_top), Inches(width),
                                Inches(y_bot - y_top))
    br.adjustments[0] = 0.04
    br.adjustments[1] = 0.5
    br.line.color.rgb = rgb(color)
    br.line.width = Pt(0.75)
    br.fill.background()
    no_effects(br)
    br.name = name
    return br


def hbracket(slide, x1, x2, y, name, color="accent", tick=0.06, width=1.0):
    """Horizontal bracket (line with end ticks) spanning a range on an axis, e.g. 'the unsecured notes'."""
    line(slide, x1, y, x2, y, color, width, name=name)
    for xx in (x1, x2):
        line(slide, xx, y - tick, xx, y + tick, color, width, name=name + " tick")


def numbered_items(slide, x, y, w, items, heights, name="item", size=11, num_d=0.24, gap=0.10):
    """Numbered dark circles + paragraphs. items = list of paras (each a list of runs lists)."""
    for k, paras in enumerate(items):
        circ = box(slide, x, y + 0.01, num_d, num_d, fill="ink", shape=MSO_SHAPE.OVAL, name=f"{name} number {k+1}")
        tf = circ.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_tf(tf, [[(str(k + 1), {"b": True, "c": "white"})]], size, align=PP_ALIGN.CENTER)
        textbox(slide, x + num_d + 0.12, y, w - num_d - 0.12, heights[k], paras, size=size, name=f"{name} {k+1}",
                space_after=2)
        y += heights[k] + gap
    return y


def card(slide, x, y, w, title, body_paras, body_h, name="card", hdr_h=0.22, body_size=9, hdr_size=10.5,
         hdr_fill="ink", hdr_color="white", body_line="grey_mid", body_dash=None):
    """Dark rounded header bar over a white body box (entity charts, business-model boxes)."""
    hb = box(slide, x, y, w, hdr_h, fill=hdr_fill, line=None, radius=0.15, name=f"{name} header")
    hb.text_frame.margin_top = hb.text_frame.margin_bottom = 0
    hb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_tf(hb.text_frame, [[(title, {"b": True, "c": hdr_color})]], hdr_size, align=PP_ALIGN.CENTER)
    bb = box(slide, x, y + hdr_h + 0.03, w, body_h, fill="white", line=body_line, dash=body_dash,
             line_w=0.75 if body_dash is None else 1.0, radius=0.06, name=f"{name} body")
    bb.text_frame.margin_top = bb.text_frame.margin_bottom = Inches(0.02)
    bb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_tf(bb.text_frame, body_paras, body_size, align=PP_ALIGN.CENTER)
    return y + hdr_h + 0.03 + body_h


# ===================================================================== tables (textbox grid)
def grid_table(slide, x, y, col_w, header, rows, name, size=10.5, row_h=0.26, left_cols=(0,),
               bold_rows=(), accent_rows=(), rule_before=(), hair_rules=False, header_size=None,
               outline_col=None, highlight_rows=(), accent_cols=None):
    """Banker table built from textboxes + rules (crisp, fully editable, no PowerPoint table styling).
    header: list of cells; a cell is a str or a list of (text, opts) runs. rows: list of lists of cells.
    accent_rows: rows whose numeric cells are accent-colored (e.g., shortfall / impaired values).
    outline_col: column index to outline with a thin ink rounded box (the recommended / base case column).
    highlight_rows: rows shaded with THEME['band'] (e.g., the fulcrum classes). Returns the y below the table."""
    W = sum(col_w)
    xs = [x + sum(col_w[:j]) for j in range(len(col_w))]

    def cell_runs(cell, base):
        runs = cell if isinstance(cell, list) else [(cell, {})]
        return [(t, {**base, **o}) for t, o in runs]

    for j, h in enumerate(header):
        textbox(slide, xs[j], y, col_w[j], 0.2, [cell_runs(h, {"b": True})], size=header_size or size,
                align=PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.RIGHT, name=f"{name} hdr {j+1}")
    line(slide, x, y + 0.25, x + W, y + 0.25, "ink", 0.75, name=f"{name} hdr rule")
    yy = y + 0.31
    for k in highlight_rows:
        box(slide, x, yy + k * row_h - 0.02, W, row_h, fill="band", line=None, shape=MSO_SHAPE.RECTANGLE,
            name=f"{name} highlight {k+1}")
    for k, row in enumerate(rows):
        if k in rule_before:
            line(slide, x, yy - 0.03, x + W, yy - 0.03, "ink", 0.75, name=f"{name} total rule")
        for j, cell in enumerate(row):
            acc = (k in accent_rows and j not in left_cols) or (accent_cols and (k, j) in accent_cols)
            base = {"b": k in bold_rows, "c": "accent" if acc else "ink"}
            textbox(slide, xs[j], yy, col_w[j], row_h - 0.06, [cell_runs(cell, base)], size=size,
                    align=PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
                    name=f"{name} r{k+1} c{j+1}")
        yy += row_h
        if hair_rules and k < len(rows) - 1 and (k + 1) not in rule_before:
            line(slide, x, yy - 0.04, x + W, yy - 0.04, "hair", 0.5, name=f"{name} rule {k+1}")
    if outline_col is not None:
        ox = xs[outline_col] + 0.06
        box(slide, ox, y - 0.03, col_w[outline_col] - 0.02, yy - y - 0.02, fill=None, line="ink", line_w=0.75,
            radius=0.06, name=f"{name} outline col {outline_col+1}")
    return yy


# ===================================================================== native PowerPoint tables
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _ln_xml(tag, color=None, w=0.5):
    if color is None:
        return f'<a:{tag} w="0"><a:noFill/></a:{tag}>'
    return (f'<a:{tag} w="{int(w*12700)}" cap="flat" cmpd="sng" algn="ctr"><a:solidFill>'
            f'<a:srgbClr val="{c(color)}"/></a:solidFill><a:prstDash val="solid"/></a:{tag}>')


def style_cell(cell, runs, size=9, align=PP_ALIGN.LEFT, fill=None, bottom=None, top=None, bottom_w=0.5,
               top_w=0.5, indent=0.0):
    """Restyle a python-pptx table cell completely (no table-style bleed): margins, borders, fill, runs.
    Use for dense data tables (debt tables) where a real table object is wanted."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        tcPr.remove(child)
    tcPr.set("marL", str(int(Inches(0.05 + indent))))
    tcPr.set("marR", str(int(Inches(0.05))))
    tcPr.set("marT", str(int(Inches(0.01))))
    tcPr.set("marB", str(int(Inches(0.01))))
    tcPr.set("anchor", "ctr")
    xml = _ln_xml("lnL") + _ln_xml("lnR") + _ln_xml("lnT", top, top_w) + _ln_xml("lnB", bottom, bottom_w)
    for el in etree.fromstring(f'<x xmlns:a="{A_NS}">{xml}</x>'):
        tcPr.append(el)
    if fill:
        f = etree.SubElement(tcPr, qn("a:solidFill"))
        etree.SubElement(f, qn("a:srgbClr")).set("val", c(fill))
    else:
        etree.SubElement(tcPr, qn("a:noFill"))
    tfc = cell.text_frame
    tfc.word_wrap = True
    p = tfc.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = align
    for text, o in runs:
        if not text:
            continue
        r = p.add_run()
        r.text = text
        style_run(r, o.get("size", size), o.get("b", False), o.get("i", False), o.get("c", "ink"),
                  o.get("sup", False))


# ===================================================================== slides
def _clear_footer(ph):
    tf = ph.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    return tf


def new_content_slide(prs, title, takeaway, notes=None, sources=None, template_index=None):
    """Add a slide on the template content slide's layout: title, one-line takeaway, footer copied from the
    template slide with Notes / Sources lines. Returns the slide (appended at the end; use move_slide)."""
    ref = prs.slides[THEME["template_slide"] if template_index is None else template_index]
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx in THEME["drop_idx"]:
            ph._element.getparent().remove(ph._element)
    s.placeholders[THEME["title_idx"]].text_frame.text = title
    s.placeholders[THEME["takeaway_idx"]].text_frame.text = takeaway
    tree = s.shapes._spTree
    for sp in ref.shapes._spTree.iterchildren(qn("p:sp")):
        ph = sp.find(".//" + qn("p:ph"))
        if ph is not None and ph.get("type") in THEME["copy_ph_types"]:
            tree.append(copy.deepcopy(sp))
    ftr = [x for x in s.placeholders if x.placeholder_format.idx == THEME["footer_idx"]]
    if ftr and (notes or sources):
        f = ftr[0]
        f.top = Inches(THEME["footer_top"])
        tf = _clear_footer(f)
        first = True
        for t in (notes, sources):
            if not t:
                continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            r = p.add_run()
            r.text = t
            r.font.size = Pt(THEME["footer_size"])
            first = False
    return s


def move_slide(prs, slide, index):
    """Move a slide to 0-based position `index`."""
    lst = prs.slides._sldIdLst
    target = next(sid for sid in list(lst) if prs.part.related_part(sid.get(qn("r:id"))) is slide.part)
    lst.remove(target)
    lst.insert(index, target)


def drop_slide(prs, index):
    """Remove the slide at 0-based `index` (its part is dropped on save)."""
    lst = prs.slides._sldIdLst
    sid = list(lst)[index]
    prs.part.drop_rel(sid.get(qn("r:id")))
    lst.remove(sid)


def copy_picture(src_slide, pic_name, dst_slide, cx, cy, scale=1.0, new_name=None):
    """Copy a picture (e.g., a logo the user already placed elsewhere in the deck) onto another slide,
    centered at (cx, cy). Keeps crops, recolor and SVG fallbacks; re-links the image parts."""
    pics = {e.find(".//" + qn("p:cNvPr")).get("name"): e for e in src_slide.shapes._spTree.iter(qn("p:pic"))}
    el = copy.deepcopy(pics[pic_name])
    for node in el.iter():
        rid = node.get(qn("r:embed"))
        if rid:
            node.set(qn("r:embed"), dst_slide.part.relate_to(src_slide.part.related_part(rid), RT.IMAGE))
    nv = el.find(".//" + qn("p:cNvPr"))
    tree = dst_slide.shapes._spTree
    nv.set("id", str(max(int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))) + 1))
    nv.set("name", new_name or f"Logo {pic_name}")
    for ext in nv.findall(qn("a:extLst")):
        nv.remove(ext)
    xf = el.find(".//" + qn("a:xfrm"))
    ext = xf.find(qn("a:ext"))
    w, h = int(int(ext.get("cx")) * scale), int(int(ext.get("cy")) * scale)
    ext.set("cx", str(w))
    ext.set("cy", str(h))
    xf.find(qn("a:off")).set("x", str(int(Inches(cx)) - w // 2))
    xf.find(qn("a:off")).set("y", str(int(Inches(cy)) - h // 2))
    tree.append(el)
    return el


def add_image(slide, path, x, y, w=None, h=None, name=None):
    """Place an image keeping its aspect ratio; give w or h. Returns (right_x, bottom_y)."""
    from PIL import Image
    W, H = Image.open(path).size
    if w is None:
        w = h * W / H
    if h is None:
        h = w * H / W
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        pic.name = name
    return x + w, y + h


# ===================================================================== native charts (XML-level control)
def col(i):
    return chr(ord("B") + i)


def txpr(size, color, bold=False):
    b = ' b="1"' if bold else ' b="0"'
    return (f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="{int(size*100)}"{b}>'
            f'<a:solidFill><a:srgbClr val="{c(color)}"/></a:solidFill>'
            f'<a:latin typeface="{THEME["font"]}"/></a:defRPr></a:pPr><a:endParaRPr lang="en-US"/></a:p></c:txPr>')


def esc(s):
    """XML-escape text that goes into chart XML (series names, categories). '&' breaks the parse otherwise."""
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def str_cache(vals):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{esc(v)}</c:v></c:pt>' for i, v in enumerate(vals))
    return f'<c:strCache><c:ptCount val="{len(vals)}"/>{pts}</c:strCache>'


def num_cache(vals, fmt="General"):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals) if v is not None)
    return f'<c:numCache><c:formatCode>{fmt}</c:formatCode><c:ptCount val="{len(vals)}"/>{pts}</c:numCache>'


def dlbls(fmt="", pos=None, size=9, color="ink", bold=True, show=True):
    if not show:
        return '<c:dLbls><c:delete val="1"/></c:dLbls>'
    pos_xml = f'<c:dLblPos val="{pos}"/>' if pos else ""
    fmt = fmt.replace("&", "&amp;").replace('"', "&quot;")
    return (f'<c:dLbls><c:numFmt formatCode="{fmt}" sourceLinked="0"/>'
            f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>{txpr(size, color, bold)}'
            f'{pos_xml}<c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/>'
            f'<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


NOLBL = '<c:dLbls><c:delete val="1"/></c:dLbls>'


def _sp_fill(fill):
    if fill is None:
        return '<a:noFill/>'
    return f'<a:solidFill><a:srgbClr val="{c(fill)}"/></a:solidFill>'


def bar_ser(i, name, cats, vals, fill, labels=NOLBL, dpt=None, outline=None, dash=None):
    """Bar/column series. dpt = {point_index: fill or (fill, outline_color, dash)} for per-point styling,
    e.g. the subject company in accent, a hollow dashed segment (fill None + outline)."""
    n = len(cats)
    pts = ""
    for idx, spec in (dpt or {}).items():
        f, o, d = (spec, None, None) if not isinstance(spec, tuple) else (spec + (None, None))[:3]
        ln = (f'<a:ln w="12700"><a:solidFill><a:srgbClr val="{c(o)}"/></a:solidFill>'
              + (f'<a:prstDash val="{d}"/>' if d else "") + '</a:ln>') if o else '<a:ln><a:noFill/></a:ln>'
        pts += (f'<c:dPt><c:idx val="{idx}"/><c:invertIfNegative val="0"/><c:bubble3D val="0"/>'
                f'<c:spPr>{_sp_fill(f)}{ln}</c:spPr></c:dPt>')
    ln = (f'<a:ln w="12700"><a:solidFill><a:srgbClr val="{c(outline)}"/></a:solidFill>'
          + (f'<a:prstDash val="{dash}"/>' if dash else "") + '</a:ln>') if outline else '<a:ln><a:noFill/></a:ln>'
    return (f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/>'
            f'<c:tx><c:strRef><c:f>Sheet1!${col(i)}$1</c:f>{str_cache([name])}</c:strRef></c:tx>'
            f'<c:spPr>{_sp_fill(fill)}{ln}</c:spPr><c:invertIfNegative val="0"/>{pts}{labels}'
            f'<c:cat><c:strRef><c:f>Sheet1!$A$2:$A${n+1}</c:f>{str_cache(cats)}</c:strRef></c:cat>'
            f'<c:val><c:numRef><c:f>Sheet1!${col(i)}$2:${col(i)}${n+1}</c:f>{num_cache(vals)}</c:numRef></c:val>'
            f'</c:ser>')


def nofill_ser(i, name, cats, vals):
    """Invisible series: the 'base' that floats stacked bars (waterfalls, football fields)."""
    return bar_ser(i, name, cats, vals, None)


def line_ser(i, name, cats, vals, color, labels=NOLBL, marker="circle", msize=6, mfill="white", width=1.5,
             dash=None, show_line=True):
    n = len(cats)
    d = f'<a:prstDash val="{dash}"/>' if dash else ""
    ln = (f'<a:ln w="{int(width*12700)}" cap="rnd"><a:solidFill><a:srgbClr val="{c(color)}"/></a:solidFill>{d}</a:ln>'
          if show_line else '<a:ln><a:noFill/></a:ln>')
    mk = (f'<c:marker><c:symbol val="{marker}"/><c:size val="{msize}"/><c:spPr>{_sp_fill(mfill)}'
          f'<a:ln w="15875"><a:solidFill><a:srgbClr val="{c(color)}"/></a:solidFill></a:ln></c:spPr></c:marker>'
          if marker != "none" else '<c:marker><c:symbol val="none"/></c:marker>')
    return (f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/>'
            f'<c:tx><c:strRef><c:f>Sheet1!${col(i)}$1</c:f>{str_cache([name])}</c:strRef></c:tx>'
            f'<c:spPr>{ln}</c:spPr>{mk}{labels}'
            f'<c:cat><c:strRef><c:f>Sheet1!$A$2:$A${n+1}</c:f>{str_cache(cats)}</c:strRef></c:cat>'
            f'<c:val><c:numRef><c:f>Sheet1!${col(i)}$2:${col(i)}${n+1}</c:f>{num_cache(vals)}</c:numRef></c:val>'
            f'<c:smooth val="0"/></c:ser>')


def cat_ax(ax_id, cross_id, delete=True, reverse=False, pos="b"):
    o = "maxMin" if reverse else "minMax"
    return (f'<c:catAx><c:axId val="{ax_id}"/><c:scaling><c:orientation val="{o}"/></c:scaling>'
            f'<c:delete val="{1 if delete else 0}"/><c:axPos val="{pos}"/><c:majorTickMark val="none"/>'
            f'<c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>{txpr(10, "grey_txt")}'
            f'<c:crossAx val="{cross_id}"/><c:crosses val="autoZero"/>'
            f'<c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>')


def val_ax(ax_id, cross_id, vmin, vmax, delete=True, pos="l", crosses="autoZero"):
    return (f'<c:valAx><c:axId val="{ax_id}"/><c:scaling><c:orientation val="minMax"/>'
            f'<c:max val="{vmax}"/><c:min val="{vmin}"/></c:scaling><c:delete val="{1 if delete else 0}"/>'
            f'<c:axPos val="{pos}"/><c:numFmt formatCode="General" sourceLinked="0"/>'
            f'<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
            f'<c:crossAx val="{cross_id}"/><c:crosses val="{crosses}"/><c:crossBetween val="between"/></c:valAx>')


def bar_chart_xml(series_xml, kind="col", grouping="clustered", gap=60, overlap=None, ax=(3001, 3002),
                  vmin=0, vmax=100, reverse_cats=None):
    """Body for a bar/column chart with hidden axes. kind 'bar' = horizontal (categories read top-down)."""
    if reverse_cats is None:
        reverse_cats = kind == "bar"
    ov = f'<c:overlap val="{overlap}"/>' if overlap is not None else (
        '<c:overlap val="100"/>' if grouping in ("stacked", "percentStacked") else "")
    return (f'<c:barChart><c:barDir val="{kind}"/><c:grouping val="{grouping}"/><c:varyColors val="0"/>'
            + "".join(series_xml)
            + f'<c:gapWidth val="{gap}"/>{ov}<c:axId val="{ax[0]}"/><c:axId val="{ax[1]}"/></c:barChart>'
            + cat_ax(ax[0], ax[1], reverse=reverse_cats, pos="l" if kind == "bar" else "b")
            + val_ax(ax[1], ax[0], vmin, vmax, pos="b" if kind == "bar" else "l"))


def chart_space(body, rid, inner=(0, 0, 1, 1)):
    """Wrap a plot body. inner = manual inner-plot layout as fractions; (0,0,1,1) makes the plot fill the
    graphic frame exactly, so data -> inch mapping is linear and labels/annotations can be placed by hand."""
    x, y, w, h = inner
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<c:chartSpace xmlns:c="{C_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}">'
            f'<c:date1904 val="0"/><c:lang val="en-US"/><c:roundedCorners val="0"/>'
            f'<c:chart><c:autoTitleDeleted val="1"/><c:plotArea>'
            f'<c:layout><c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
            f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{h}"/></c:manualLayout></c:layout>'
            f'{body}<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr></c:plotArea>'
            f'<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/></c:chart>'
            f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>{txpr(10, "grey_txt")}'
            f'<c:externalData r:id="{rid}"><c:autoUpdate val="0"/></c:externalData></c:chartSpace>')


def add_chart(slide, x, y, w, h, cats, series, body_builder, name, inner=(0, 0, 1, 1)):
    """Create a native, editable chart (embedded workbook holds every series), then replace its XML.
    series = [(name, values), ...] in the same order as the series XML. body_builder() -> plot body XML."""
    cd = CategoryChartData()
    cd.categories = cats
    for s_name, vals in series:
        cd.add_series(s_name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    part = gf.chart_part
    rid = part._element.find(qn("c:externalData")).get(qn("r:id"))
    part._element = parse_xml(chart_space(body_builder(), rid, inner).encode("utf-8"))
    return gf


class Plot:
    """Linear data -> inch mapping for a chart drawn with chart_space(inner=(0,0,1,1)) and hidden axes.
    horizontal=False: categories along x (columns). horizontal=True: categories down y (bars)."""

    def __init__(self, x, y, w, h, n, vmin, vmax, horizontal=False):
        self.x, self.y, self.w, self.h, self.n = x, y, w, h, n
        self.vmin, self.vmax, self.horizontal = vmin, vmax, horizontal
        self.slot = (h if horizontal else w) / n

    def center(self, k):
        return (self.y if self.horizontal else self.x) + self.slot * (k + 0.5)

    def v(self, val):
        """Value -> inch coordinate (y for columns, x for bars)."""
        f = (val - self.vmin) / (self.vmax - self.vmin)
        return self.x + self.w * f if self.horizontal else self.y + self.h * (1 - f)

    def bar_width(self, gap):
        return self.slot / (1 + gap / 100)


def chart(slide, plot, cats, series, series_xml, name, kind=None, grouping="clustered", gap=60, overlap=None):
    """Convenience: native chart on a Plot geometry. series = [(name, vals)], series_xml = list of ser XML."""
    kind = kind or ("bar" if plot.horizontal else "col")
    return add_chart(slide, plot.x, plot.y, plot.w, plot.h, cats, series,
                     lambda: bar_chart_xml(series_xml, kind, grouping, gap, overlap, vmin=plot.vmin, vmax=plot.vmax),
                     name)


__all__ = [n for n in globals() if not n.startswith("_")]
