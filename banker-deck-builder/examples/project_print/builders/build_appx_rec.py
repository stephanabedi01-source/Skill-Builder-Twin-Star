"""Append two versions of the appendix page 'Recovery Detail by Tranche' (pick one):
  A: model-output version (tranche table + waterfall assumptions)
  B: same tranche table + sensitivity to the two contested claims (pension, Holdings intercompany double-dip)

Usage: python3 build_appx_rec.py <input.pptx> <output.pptx>
Numbers: model/v16_verify.xlsx ($XRX Recovery Waterfall); toggle cases recalculated on copies in model/tog/.
"""
import sys

_argv = list(sys.argv)
sys.argv = sys.argv[:1]
from openpyxl import load_workbook  # noqa: E402
from pp_helpers import *  # noqa: E402,F401,F403
from pp_helpers import textbox, section_header  # noqa: E402

SRC, OUT = _argv[1], _argv[2]
HAIR = "D1D1D1"
GREY_FILL = "D6D6D6"
BAND = "E9E9E9"
L0, R1 = 0.23, 13.11
LP1, RP0, ROW1 = 8.33, 8.62, 1.30

W = load_workbook("model/v16_verify.xlsx", data_only=True)["$XRX Recovery Waterfall"]
TOG = {n: load_workbook(f"model/tog/{n}.xlsx", data_only=True)["$XRX Recovery Waterfall"]
       for n in ("both_on", "pension_off", "dd_off", "both_off")}

prs = Presentation(SRC)
ref = prs.slides[2]


def new_slide(takeaway, notes, sources):
    s = prs.slides.add_slide(ref.slide_layout)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx == 20:
            ph._element.getparent().remove(ph._element)
    s.placeholders[21].text_frame.text = "Appendix: Recovery Detail by Tranche"
    s.placeholders[16].text_frame.text = takeaway
    tree = s.shapes._spTree
    for sp in ref.shapes._spTree.iterchildren(qn("p:sp")):
        ph = sp.find(".//" + qn("p:ph"))
        if ph is not None and ph.get("type") in ("sldNum", "ftr"):
            tree.append(copy.deepcopy(sp))
    f = [x for x in s.placeholders if x.placeholder_format.idx == 18][0]
    f.top = Inches(6.93)
    tf = f.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    for i, t in enumerate((notes, sources)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = t
        r.font.size = Pt(10)
    return s


def callout(s, x, y, w, h, runs, name, size=10.5):
    c = box(s, x, y, w, h, fill=GREY_FILL, line=RED, dash=MSO_LINE_DASH_STYLE.ROUND_DOT, line_w=1.0,
            radius=0.08, name=name)
    c.text_frame.margin_left = c.text_frame.margin_right = Inches(0.08)
    c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_tf(c.text_frame, [runs], size, align=PP_ALIGN.CENTER)
    return c


def money(v):
    return "n.a." if v is None else (f"${v:,.0f}" if v >= 0.5 else "$0")


def pct(v):
    return "n.a." if v is None else f"{v:.0%}"


# ---------------------------------------------------------------- tranche table (shared by both versions)
TRANCHES = [  # (label, row, group)
    ("IPCo JV term loans (incl. premium)", 67, "sec"), ("IPCo Class A units", 68, "sec"),
    ("ABL revolver (60% drawn)", 66, "sec"), ("1L Term Loan B due 2029", 69, "sec"),
    ("1L 10.25% notes due 2030", 70, "sec"), ("2L 13.50% notes due 2031", 71, "sec"),
    ("Holdings 5.50% notes due 2028", 72, "hold"), ("Holdings 8.875% notes due 2029", 73, "hold"),
    ("Holdings 3.75% converts due 2030", 74, "hold"), ("Holdings 13.00% step-up notes due 2030", 75, "hold"),
    ("Xerox Corp. 4.80% notes due 2035", 76, "leg"), ("Xerox Corp. 6.75% notes due 2039", 77, "leg"),
    ("Pension / OPEB underfunding", 78, "leg"), ("Series A preferred", 79, "eq"), ("Common equity", 80, "eq")]


def tranche_table(s, x0, y0, width):
    CW = [2.62] + [0.78] + [0.78, 0.66] * 3
    scale = width / sum(CW)
    CW = [c * scale for c in CW]
    xs = [x0 + sum(CW[:j]) for j in range(len(CW))]
    # group headers
    for k, (lbl, col) in enumerate((("Low EV", 2), ("Mid EV", 4), ("High EV", 6))):
        textbox(s, xs[col], y0, CW[col] + CW[col + 1], 0.18, [[(f"{lbl} ${W.cell(25, 5 + 2 * k).value:,.0f}M",
                                                               {"b": True})]], size=9.5, align=PP_ALIGN.CENTER,
                name=f"T group {k + 1}")
        line(s, xs[col] + 0.05, y0 + 0.21, xs[col] + CW[col] + CW[col + 1] - 0.05, y0 + 0.21, INK, 0.5,
             name=f"T group rule {k + 1}")
    hdr = ["($M)", "Claim", "Recovery", "%", "Recovery", "%", "Recovery", "%"]
    for j, h in enumerate(hdr):
        textbox(s, xs[j], y0 + 0.25, CW[j], 0.18, [[(h, {"b": True})]], size=9.5,
                align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT, name=f"T hdr {j + 1}")
    line(s, x0, y0 + 0.46, x0 + width, y0 + 0.46, INK, 0.75, name="T hdr rule")
    RH = 0.215
    y = y0 + 0.50
    # admin row first (from distributable value block)
    rows = [("Admin and priority claims", 125.0, [125.0, 1.0] * 3, "sec")]
    for lbl, r, grp in TRANCHES:
        vals = []
        for c in "EFGHIJ":
            vals.append(W[f"{c}{r}"].value)
        rows.append((lbl, W[f"D{r}"].value, vals, grp))
    prev = None
    for k, (lbl, claim, vals, grp) in enumerate(rows):
        if prev is not None and grp != prev:
            line(s, x0, y - 0.01, x0 + width, y - 0.01, GREY_MID, 0.5, name=f"T grp rule {k}")
        hot = grp in ("hold", "leg")
        cells = [lbl, money(claim)] + [money(v) if i % 2 == 0 else pct(v) for i, v in enumerate(vals)]
        for j, t in enumerate(cells):
            o = {"c": RED if (hot and j in (3, 5, 7)) or (grp == "eq" and j in (3, 5, 7)) else INK,
                 "b": j in (3, 5, 7)}
            textbox(s, xs[j], y, CW[j], RH - 0.02, [[(t, o)]], size=9.5,
                    align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
                    name=f"T r{k + 1} c{j + 1}")
        y += RH
        prev = grp
    # total row
    line(s, x0, y + 0.01, x0 + width, y + 0.01, INK, 0.75, name="T total rule")
    tot = ["Total distributed", money(W["D81"].value + 125)]
    for c in "EG I".replace(" ", ""):
        pass
    tot += [money(W["E81"].value + 125), pct((W["E81"].value + 125) / (W["D81"].value + 125)),
            money(W["G81"].value + 125), pct((W["G81"].value + 125) / (W["D81"].value + 125)),
            money(W["I81"].value + 125), pct((W["I81"].value + 125) / (W["D81"].value + 125))]
    y += 0.04
    for j, t in enumerate(tot):
        textbox(s, xs[j], y, CW[j], RH - 0.02, [[(t, {"b": True})]], size=9.5,
                align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
                name=f"T total c{j + 1}")
    return y + RH


def small_table(s, x0, y0, cw, hdr, rows, name, bold_last=False, size=10, rh=0.25, hot_rows=()):
    for j, h in enumerate(hdr):
        textbox(s, x0 + sum(cw[:j]), y0, cw[j], 0.2, [[(h, {"b": True})]], size=size,
                align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT, name=f"{name} hdr {j + 1}")
    w = sum(cw)
    line(s, x0, y0 + 0.24, x0 + w, y0 + 0.24, INK, 0.75, name=f"{name} hdr rule")
    y = y0 + 0.29
    for k, row in enumerate(rows):
        last = bold_last and k == len(rows) - 1
        if last:
            line(s, x0, y - 0.02, x0 + w, y - 0.02, INK, 0.75, name=f"{name} total rule")
        for j, t in enumerate(row):
            runs = t if isinstance(t, list) else [(t, {})]
            runs = [(a, {"b": last or (k in hot_rows and j > 0), **o}) for a, o in runs]
            textbox(s, x0 + sum(cw[:j]), y, cw[j], rh - 0.03, [runs], size=size,
                    align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
                    name=f"{name} r{k + 1} c{j + 1}")
        y += rh
        if k < len(rows) - 1 and not (bold_last and k == len(rows) - 2):
            line(s, x0, y - 0.02, x0 + w, y - 0.02, HAIR, 0.5, name=f"{name} rule {k + 1}")
    return y


NOTES = ("Notes: Claims include 3 months of accrued interest; ABL assumed 60% drawn at filing; IPCo JV repaid first from "
         "its 17% share of EV incl. 3.875% premium; Holdings notes claim on the Xerox Corp. guarantee and the "
         "intercompany loan, capped at par")
SOURCES = "Sources: Xerox Q2'26 10-Q; indentures; Project Print model ($XRX Recovery Waterfall)"

# ================================================================= VERSION A: model output
sA = new_slide("At every plan value, secured claims are paid in full, value breaks in the unsecured notes and equity "
               "recovers nothing", NOTES, SOURCES)
section_header(sA, L0, ROW1, LP1 - L0, "Recoveries by Tranche, Low / Mid / High Plan Value", "A header")
tranche_table(sA, L0, 1.70, LP1 - L0)
section_header(sA, RP0, ROW1, R1 - RP0, "Distributable Value ($M)", "B header")
dv = [["Enterprise value"] + [money(W.cell(27, c).value) for c in (5, 7, 9)],
      ["(-) Admin and priority claims"] + ["(125)"] * 3,
      ["Net distributable value"] + [money(W.cell(29, c).value) for c in (5, 7, 9)],
      ["EV / LTM Adj. EBITDA"] + [f"{W.cell(25, c).value / 614:.1f}x" for c in (5, 7, 9)]]
small_table(sA, RP0, 1.72, [2.02, 0.82, 0.82, 0.83], ["", "Low", "Mid", "High"], dv, "B dv", hot_rows=(2,))
section_header(sA, RP0, 3.10, R1 - RP0, "Waterfall Assumptions", "C header")
asm = [["Accrued interest included in claims", "3 months"],
       ["ABL drawn at filing (% of availability)", "60%"],
       ["Administrative and priority claims", "$125M"],
       ["IPCo JV share of EV (brand IP)", "17%"],
       ["IPCo prepayment premium (after 2/17/27)", "3.875%"],
       ["Pension / OPEB claim (pari passu)", "$1,175M, included"],
       ["Holdings intercompany loan claim", "Included"],
       ["Fulcrum security (all cases)", [("Unsecured notes", {"c": RED, "b": True})]]]
small_table(sA, RP0, 3.52, [2.85, 1.64], ["Assumption", "Input"], asm, "C asm", rh=0.27)

# ================================================================= VERSION B: with contested-claim sensitivity
sB = new_slide("However the contested claims are treated, Holdings recover 27c to 46c at mid value: the prepack's 52c "
               "beats every case", NOTES + "; sensitivity cases recalculated in the same model with the toggles off",
               SOURCES)
section_header(sB, L0, ROW1, LP1 - L0, "Recoveries by Tranche, Low / Mid / High Plan Value", "A header")
tranche_table(sB, L0, 1.70, LP1 - L0)
section_header(sB, RP0, ROW1, R1 - RP0, "Sensitivity to the Contested Claims", "B header")
textbox(sB, RP0, 1.66, R1 - RP0, 0.2, [[("Recovery at mid EV (low to high)", {"i": True, "c": GREY_TXT})]], size=10,
        name="B caption")
cases = [("Both claims allowed (base)", "both_on"), ("Pension claim excluded", "pension_off"),
         ("No Holdings double-dip", "dd_off"), ("Both excluded", "both_off")]


def rng(ws, r):
    lo, mid, hi = (ws[f"{c}{r}"].value for c in "FHJ")
    return [(f"{mid:.0%}", {"b": True}), (f" ({lo * 100:.0f}-{hi * 100:.0f})", {"c": GREY_TXT, "size": 9})]


srows = []
for lbl, n in cases:
    ws = TOG[n]
    pen = ws["H78"].value
    srows.append([lbl, rng(ws, 72), rng(ws, 76), "0%" if not pen else f"{pen:.0%}"])
SCW = [1.62, 1.04, 1.04, 0.79]
yS = small_table(sB, RP0, 1.92, SCW, ["Case", "Holdings", "Legacy", "Pension"], srows,
                 "B sens", rh=0.30)
mkt = [["Market price, 10/14/26", [("46c", {"b": True})], [("32c", {"b": True})], ""],
       ["Prepack plan", [("52c", {"b": True, "c": RED})], [("26c", {"b": True, "c": RED})], "100%"]]
yR = small_table(sB, RP0, yS + 0.08, SCW, ["Reference", "", "", ""], mkt, "B ref", rh=0.28)
callout(sB, RP0, yR + 0.10, R1 - RP0, 0.78,
        [("Excluding the pension claim reproduces the ", {}), ("46c market price", {"b": True}),
         (" for the Holdings notes. The prepack's ", {}), ("52c", {"b": True, "c": RED}),
         (" beats the Holdings mid recovery in every case; the double-dip is what justifies paying Holdings ~2x the "
          "legacy notes", {})], "B callout", size=10)
section_header(sB, RP0, yR + 1.05, R1 - RP0, "Why These Two Claims Are Contested", "C header")
textbox(sB, RP0, yR + 1.43, R1 - RP0, 0.62,
        [[("Pension: ", {"b": True}), ("PBGC claim size and priority are litigated in most cases", {"c": GREY_TXT})],
         [("Double-dip: ", {"b": True}), ("the intercompany loan can be challenged as equity or disallowed",
                                          {"c": GREY_TXT})]], size=10, space_after=3, name="C text")

prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
