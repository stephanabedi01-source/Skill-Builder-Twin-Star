"""Remove personal pronouns from slides 6-8. Usage: python3 fix_pronouns.py <in.pptx> <out.pptx>"""
import sys
from pptx import Presentation

FIX = {
    (6, "Business plan conclusion"): [("the basis for our valuation", "the basis for the valuation")],
    (7, "C callout"): [(": we value the business, not the share price, because each class of claims is paid from EV",
                        ": each class of claims is paid from enterprise value, not from the share price")],
    (8, "C callout"): [("We apply ", "The selected "), (", below the peer band: Xerox carries ",
                                                        " sits below the peer band: Xerox carries ")],
}
prs = Presentation(sys.argv[1])
for (n, name), pairs in FIX.items():
    sh = next(s for s in prs.slides[n - 1].shapes if s.name == name)
    for old, new in pairs:
        hit = [r for p in sh.text_frame.paragraphs for r in p.runs if old in r.text]
        if not hit:      # already fixed at build time
            continue
        hit[0].text = hit[0].text.replace(old, new)
prs.save(sys.argv[2])
print("saved", sys.argv[2])
