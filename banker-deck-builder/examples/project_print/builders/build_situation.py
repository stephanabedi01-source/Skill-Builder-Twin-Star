"""Add the Situation Overview slide (slide 4) to Project Print deck v2.

Usage: python3 build_situation.py <input.pptx> <output.pptx>
"""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC, OUT = sys.argv[1], sys.argv[2]

# ---------------------------------------------------------------- palette
INK = "1A1A24"
RED = "D92332"      # lines, markers, pills, red text (matches slide 3)
RED_BAR = "C00000"  # bar fills (matches the Xerox bar on slide 3)
GREY_TXT = "595959"
GREY_MID = "7F8590"
GREY_LT = "BFC3C9"
WHITE = "FFFFFF"
FONT = "Garamond"


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- text helpers
def style_run(run, size, bold=False, italic=False, color=INK, sup=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    if sup:
        run._r.get_or_add_rPr().set("baseline", "30000")


def fill_tf(tf, paras, size, align=PP_ALIGN.LEFT, space_after=0, line_spacing=None):
    """paras: list of paragraphs; each paragraph is a list of (text, opts) runs."""
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, o in runs:
            r = p.add_run()
            r.text = text
            style_run(r, o.get("size", size), o.get("b", False), o.get("i", False),
                      o.get("c", INK), o.get("sup", False))


def textbox(slide, x, y, w, h, paras, size=11, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, name=None, space_after=0, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    fill_tf(tf, paras, size, align, space_after)
    if name:
        tb.name = name
    return tb


def no_effects(shape):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    spPr.append(etree.SubElement(spPr, qn("a:effectLst")))


def box(slide, x, y, w, h, fill=WHITE, line=None, dash=None, line_w=0.75,
        radius=None, name=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = dash
    no_effects(s)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.word_wrap = True
    if name:
        s.name = name
    return s


def line(slide, x1, y1, x2, y2, color=INK, width=1.0, dash=None, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if name:
        c.name = name
    return c


def section_header(slide, x, y, w, text, name):
    textbox(slide, x, y, w, 0.24, [[(text, {"b": True})]], size=14, name=name)
    line(slide, x, y + 0.29, x + w, y + 0.29, INK, 1.0, name=name + " rule")


# ---------------------------------------------------------------- chart XML
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def col(i):
    return chr(ord("B") + i)


def txpr(size, color, bold=False):
    b = ' b="1"' if bold else ' b="0"'
    return (f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="{int(size*100)}"{b}>'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="{FONT}"/></a:defRPr></a:pPr><a:endParaRPr lang="en-US"/></a:p></c:txPr>')


def str_cache(vals):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals))
    return f'<c:strCache><c:ptCount val="{len(vals)}"/>{pts}</c:strCache>'


def num_cache(vals, fmt="General"):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals)
                  if v is not None)
    return f'<c:numCache><c:formatCode>{fmt}</c:formatCode><c:ptCount val="{len(vals)}"/>{pts}</c:numCache>'


def dlbls(fmt, pos, size, color, bold=True, show=True, deleted_idx=()):
    if not show:
        return '<c:dLbls><c:delete val="1"/></c:dLbls>'
    dels = "".join(f'<c:dLbl><c:idx val="{i}"/><c:delete val="1"/></c:dLbl>' for i in deleted_idx)
    pos_xml = f'<c:dLblPos val="{pos}"/>' if pos else ""
    fmt = fmt.replace('&', '&amp;').replace('"', '&quot;')
    return (f'<c:dLbls>{dels}<c:numFmt formatCode="{fmt}" sourceLinked="0"/>'
            f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>{txpr(size, color, bold)}'
            f'{pos_xml}<c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/>'
            f'<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


def bar_ser(i, name, cats, vals, fill, labels, dpt_fills=None):
    n = len(cats)
    dpts = ""
    for idx, f in (dpt_fills or {}).items():
        dpts += (f'<c:dPt><c:idx val="{idx}"/><c:invertIfNegative val="0"/><c:bubble3D val="0"/>'
                 f'<c:spPr><a:solidFill><a:srgbClr val="{f}"/></a:solidFill></c:spPr></c:dPt>')
    return (f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/>'
            f'<c:tx><c:strRef><c:f>Sheet1!${col(i)}$1</c:f>{str_cache([name])}</c:strRef></c:tx>'
            f'<c:spPr><a:solidFill><a:srgbClr val="{fill}"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>'
            f'<c:invertIfNegative val="0"/>{dpts}{labels}'
            f'<c:cat><c:strRef><c:f>Sheet1!$A$2:$A${n+1}</c:f>{str_cache(cats)}</c:strRef></c:cat>'
            f'<c:val><c:numRef><c:f>Sheet1!${col(i)}$2:${col(i)}${n+1}</c:f>{num_cache(vals)}</c:numRef></c:val>'
            f'</c:ser>')


def line_ser(i, name, cats, vals, color, labels, marker="diamond", msize=6, mfill=WHITE,
             show_line=True, mline_w=1.0):
    n = len(cats)
    ln = (f'<a:ln w="19050" cap="rnd"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:ln>'
          if show_line else '<a:ln><a:noFill/></a:ln>')
    mfill_xml = (f'<a:solidFill><a:srgbClr val="{mfill}"/></a:solidFill>' if mfill
                 else '<a:noFill/>')
    return (f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/>'
            f'<c:tx><c:strRef><c:f>Sheet1!${col(i)}$1</c:f>{str_cache([name])}</c:strRef></c:tx>'
            f'<c:spPr>{ln}</c:spPr>'
            f'<c:marker><c:symbol val="{marker}"/><c:size val="{msize}"/>'
            f'<c:spPr>{mfill_xml}<a:ln w="{int(mline_w*12700)}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:ln></c:spPr></c:marker>'
            f'{labels}'
            f'<c:cat><c:strRef><c:f>Sheet1!$A$2:$A${n+1}</c:f>{str_cache(cats)}</c:strRef></c:cat>'
            f'<c:val><c:numRef><c:f>Sheet1!${col(i)}$2:${col(i)}${n+1}</c:f>{num_cache(vals)}</c:numRef></c:val>'
            f'<c:smooth val="0"/></c:ser>')


def cat_ax(ax_id, cross_id, delete=False, lbl_pos="nextTo", line_color=GREY_LT, size=10):
    if delete:
        return (f'<c:catAx><c:axId val="{ax_id}"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
                f'<c:delete val="1"/><c:axPos val="b"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
                f'<c:tickLblPos val="nextTo"/><c:crossAx val="{cross_id}"/><c:crosses val="autoZero"/>'
                f'<c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>')
    return (f'<c:catAx><c:axId val="{ax_id}"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
            f'<c:delete val="0"/><c:axPos val="b"/><c:numFmt formatCode="General" sourceLinked="0"/>'
            f'<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="{lbl_pos}"/>'
            f'<c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="{line_color}"/></a:solidFill></a:ln></c:spPr>'
            f'{txpr(size, GREY_TXT)}<c:crossAx val="{cross_id}"/><c:crosses val="autoZero"/>'
            f'<c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>')


def val_ax(ax_id, cross_id, vmin, vmax, pos="l", crosses="autoZero", hidden_visible=False):
    if hidden_visible:
        return (f'<c:valAx><c:axId val="{ax_id}"/><c:scaling><c:orientation val="minMax"/>'
                f'<c:max val="{vmax}"/><c:min val="{vmin}"/></c:scaling><c:delete val="0"/>'
                f'<c:axPos val="{pos}"/><c:numFmt formatCode="General" sourceLinked="0"/>'
                f'<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="none"/>'
                f'<c:spPr><a:ln><a:noFill/></a:ln></c:spPr>'
                f'<c:crossAx val="{cross_id}"/><c:crosses val="{crosses}"/><c:crossBetween val="between"/></c:valAx>')
    return (f'<c:valAx><c:axId val="{ax_id}"/><c:scaling><c:orientation val="minMax"/>'
            f'<c:max val="{vmax}"/><c:min val="{vmin}"/></c:scaling><c:delete val="1"/>'
            f'<c:axPos val="{pos}"/><c:numFmt formatCode="General" sourceLinked="0"/>'
            f'<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
            f'<c:crossAx val="{cross_id}"/><c:crosses val="{crosses}"/><c:crossBetween val="between"/></c:valAx>')


def chart_space(plot_inner, plot_body, rid):
    x, y, w, h = plot_inner
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<c:chartSpace xmlns:c="{C_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}">'
            f'<c:date1904 val="0"/><c:lang val="en-US"/><c:roundedCorners val="0"/>'
            f'<c:chart><c:autoTitleDeleted val="1"/><c:plotArea>'
            f'<c:layout><c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
            f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{h}"/></c:manualLayout></c:layout>'
            f'{plot_body}<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr></c:plotArea>'
            f'<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/></c:chart>'
            f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
            f'{txpr(10, GREY_TXT)}'
            f'<c:externalData r:id="{rid}"><c:autoUpdate val="0"/></c:externalData></c:chartSpace>')


def add_chart(slide, x, y, w, h, cats, series, xml_builder, name):
    """Create a native chart (embedded workbook holds every series), then replace its XML."""
    cd = CategoryChartData()
    cd.categories = cats
    for s_name, vals in series:
        cd.add_series(s_name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
    gf.name = name
    part = gf.chart_part
    rid = part._element.find(qn("c:externalData")).get(qn("r:id"))
    part._element = parse_xml(xml_builder(rid).encode("utf-8"))
    return gf


# ---------------------------------------------------------------- open deck, add slide
prs = Presentation(SRC)
src_slide = prs.slides[2]          # Company Snapshot, used for footer/slide-number XML
layout = src_slide.slide_layout
slide = prs.slides.add_slide(layout)

# remove the optional section-header placeholder (idx 20), as on slide 3
for ph in list(slide.placeholders):
    if ph.placeholder_format.idx == 20:
        ph._element.getparent().remove(ph._element)

slide.placeholders[21].text_frame.text = "Situation Overview"
slide.placeholders[16].text_frame.text = (
    "Debt-funded M&A and negative core cash flow have left Xerox at ~7x leverage "
    "with $1.9B due in 2028 and 2029")

# copy slide-number and footer placeholders from slide 3
spTree = slide.shapes._spTree
for sp in src_slide.shapes._spTree.iterchildren(qn("p:sp")):
    ph = sp.find(".//" + qn("p:ph"))
    if ph is not None and ph.get("type") in ("sldNum", "ftr"):
        spTree.append(copy.deepcopy(sp))

ftr = [s for s in slide.placeholders if s.placeholder_format.type is not None
       and s.placeholder_format.idx == 18][0]
tf = ftr.text_frame
for p in list(tf.paragraphs)[1:]:
    p._p.getparent().remove(p._p)
tf.paragraphs[0]._p.clear() if False else None
ftr.top = Inches(6.93)
p0 = tf.paragraphs[0]
for r in list(p0.runs):
    r._r.getparent().remove(r._r)


def ftr_para(p, text):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(10)


ftr_para(p0, "Notes: (1) Total debt / TTM Adj. EBITDA per Xerox (FY2025 PF for Lexmark; FY2023 per Pari Passu); "
             "(2) Excl. $105mm IEEPA tariff benefit; (3) Reported OCF less finance receivable runoff; (4) Cash and "
             "ST investments plus ABL net availability; (5) Scheduled principal per 10-Q, incl. amortization; (6) Paraphrased")
ftr_para(tf.add_paragraph(),
         "Sources: Xerox 10-K, 10-Qs, earnings slides, Aug-26 investor deck; ABL Amendment No. 2; "
         "S&P (6/16/26); Pari Passu; FactSet as of 10/14/26")

# ---------------------------------------------------------------- geometry
L0, L1 = 0.23, 6.53
R0, R1 = 6.81, 13.11
ROW1, ROW2 = 1.30, 4.13

# ================================================================= MODULE 1
section_header(slide, L0, ROW1, L1 - L0, "Path to Distress: Debt-Funded Capital Allocation",
               "M1 header")

m1_cats = ["FY2023", "FY2024", "FY2025", "Q2'26 LTM"]
m1_debt = [3277, 3399, 4247, 4223]
m1_lev = [5.5, 6.1, 6.7, 5.9]
m1_ext = [None, None, None, 6.9]
C1 = dict(x=L0, y=2.06, w=5.0, h=2.00, px=0.02, pw=0.96, py=0.155, ph=0.70,
          pmax=4600, smax=11.0)


def m1_xml(rid):
    bars = (f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/><c:varyColors val="0"/>'
            + bar_ser(0, "Total Debt ($M)", m1_cats, m1_debt, INK,
                      dlbls('\\$#,##0"M"', "outEnd", 10.5, INK, True))
            + '<c:gapWidth val="80"/><c:axId val="1001"/><c:axId val="1002"/></c:barChart>')
    lines = ('<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/>'
             + line_ser(1, "Leverage (x)", m1_cats, m1_lev, RED,
                        dlbls('0.0"x"', "b", 10, WHITE, True))
             + line_ser(2, "Leverage ex-tariff (x)", m1_cats, m1_ext, RED,
                        dlbls("", None, 10, WHITE, show=False), marker="circle", msize=8,
                        mfill=WHITE, show_line=False, mline_w=1.75)
             + '<c:marker val="1"/><c:axId val="1003"/><c:axId val="1004"/></c:lineChart>')
    axes = (cat_ax(1001, 1002) + val_ax(1002, 1001, 0, C1["pmax"])
            + cat_ax(1003, 1004, delete=True) + val_ax(1004, 1003, 0, C1["smax"], pos="r", crosses="max", hidden_visible=True))
    return chart_space((C1["px"], C1["py"], C1["pw"], C1["ph"]), bars + lines + axes, rid)


add_chart(slide, C1["x"], C1["y"], C1["w"], C1["h"], m1_cats,
          [("Total Debt ($M)", m1_debt), ("Leverage (x)", m1_lev),
           ("Leverage ex-tariff (x)", m1_ext)], m1_xml, "M1 debt and leverage chart")

# geometry helpers for module 1
p1_x0 = C1["x"] + C1["px"] * C1["w"]
p1_w = C1["pw"] * C1["w"]
p1_y0 = C1["y"] + C1["py"] * C1["h"]
p1_h = C1["ph"] * C1["h"]
p1_bot = p1_y0 + p1_h
slot1 = p1_w / 4
bar_w1 = slot1 / 1.8
centers1 = [p1_x0 + slot1 * (k + 0.5) for k in range(4)]


def y1(v):
    return p1_bot - p1_h * v / C1["pmax"]


def y1s(v):
    return p1_bot - p1_h * v / C1["smax"]


# deal tags with connectors to the bar labels
tags = [
    [("Icahn buyback: ", {"b": True}), ("$542M, funded by $500M TL", {})],
    [("ITsavvy: ", {"b": True}), ("$400M, IT diversification", {})],
    [("Lexmark: ", {"b": True}), ("$1.5B, A4 / APAC / manufacturing; debt-funded", {})],
]
TAG_W, TAG_Y, TAG_H = 1.16, 1.65, 0.58
for k, runs in enumerate(tags):
    cx = centers1[k]
    t = box(slide, cx - TAG_W / 2, TAG_Y, TAG_W, TAG_H, fill=WHITE, line=GREY_MID,
            dash=MSO_LINE_DASH_STYLE.ROUND_DOT, line_w=0.75, radius=0.18, name=f"M1 deal tag {k+1}")
    t.text_frame.margin_left = t.text_frame.margin_right = Inches(0.05)
    t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_tf(t.text_frame, [runs], 8.5, align=PP_ALIGN.CENTER)
    label_top = y1(m1_debt[k]) - 0.22
    line(slide, cx, TAG_Y + TAG_H, cx, label_top, GREY_MID, 0.75, name=f"M1 tag connector {k+1}")

# ex-tariff pill to the right of the Q2'26 bar
ext_y = y1s(6.9)
PILL_X, PILL_W, PILL_H = 5.32, 1.18, 0.26
line(slide, centers1[3] + 0.07, ext_y, PILL_X, ext_y, RED, 0.75, name="M1 ex-tariff connector")
pill = box(slide, PILL_X, ext_y - PILL_H / 2, PILL_W, PILL_H, fill=WHITE, line=RED,
           dash=MSO_LINE_DASH_STYLE.ROUND_DOT, line_w=1.0, radius=0.5, name="M1 ex-tariff pill")
pill.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
fill_tf(pill.text_frame, [[("6.9x", {"b": True, "c": RED}), (" ex-tariff", {}), ("2", {"sup": True})]],
        10, align=PP_ALIGN.CENTER)

# legend in the right strip
LG_X, LG_Y = 5.36, 3.30
sq = box(slide, LG_X, LG_Y + 0.035, 0.11, 0.11, fill=INK, shape=MSO_SHAPE.RECTANGLE, name="M1 legend debt key")
textbox(slide, LG_X + 0.2, LG_Y, 1.0, 0.18, [[("Total debt ($M)", {"c": GREY_TXT})]], size=9,
        name="M1 legend debt")
line(slide, LG_X - 0.03, LG_Y + 0.31, LG_X + 0.14, LG_Y + 0.31, RED, 1.5, name="M1 legend lev line")
dm = box(slide, LG_X + 0.015, LG_Y + 0.26, 0.09, 0.1, fill=WHITE, line=RED, line_w=0.75,
         shape=MSO_SHAPE.DIAMOND, name="M1 legend lev marker")
textbox(slide, LG_X + 0.2, LG_Y + 0.22, 1.0, 0.18,
        [[("Leverage (x)", {"c": GREY_TXT}), ("1", {"c": GREY_TXT, "sup": True})]], size=9,
        name="M1 legend leverage")

# ================================================================= MODULE 2
section_header(slide, R0, ROW1, R1 - R0, "Quality of Cash Flow", "M2 header")

m2_cats = ["FY2024", "FY2025"]
m2_rep = [511, 224]
m2_core = [-152, -265]
C2 = dict(x=R0, y=1.88, w=4.5, h=2.17, px=0.02, pw=0.8613, py=0.10, ph=0.772,
          vmin=-420, vmax=680, gap=310)


def m2_xml(rid):
    body = (f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/><c:varyColors val="0"/>'
            + bar_ser(0, "Reported OCF ($M)", m2_cats, m2_rep, GREY_LT,
                      dlbls('\\$#,##0"M";\\$(#,##0)"M"', "outEnd", 10.5, INK, True))
            + bar_ser(1, "Core OCF ex. receivable runoff ($M)", m2_cats, m2_core, RED_BAR,
                      dlbls('\\$#,##0"M";\\$(#,##0)"M"', "outEnd", 10.5, RED, True))
            + f'<c:gapWidth val="{C2["gap"]}"/><c:overlap val="0"/><c:axId val="2001"/><c:axId val="2002"/></c:barChart>'
            + cat_ax(2001, 2002, lbl_pos="low", line_color=GREY_MID)
            + val_ax(2002, 2001, C2["vmin"], C2["vmax"]))
    return chart_space((C2["px"], C2["py"], C2["pw"], C2["ph"]), body, rid)


add_chart(slide, C2["x"], C2["y"], C2["w"], C2["h"], m2_cats,
          [("Reported OCF ($M)", m2_rep), ("Core OCF ex. receivable runoff ($M)", m2_core)],
          m2_xml, "M2 operating cash flow chart")

p2_x0 = C2["x"] + C2["px"] * C2["w"]
p2_w = C2["pw"] * C2["w"]
p2_y0 = C2["y"] + C2["py"] * C2["h"]
p2_h = C2["ph"] * C2["h"]
slot2 = p2_w / 2
bar_w2 = slot2 / (2 + C2["gap"] / 100)
margin2 = bar_w2 * C2["gap"] / 200


def y2(v):
    return p2_y0 + p2_h * (C2["vmax"] - v) / (C2["vmax"] - C2["vmin"])


# legend
LG2_Y = 1.67
box(slide, R0 + 0.02, LG2_Y + 0.035, 0.11, 0.11, fill=GREY_LT, shape=MSO_SHAPE.RECTANGLE, name="M2 legend rep key")
textbox(slide, R0 + 0.2, LG2_Y, 1.9, 0.18, [[("Reported operating cash flow", {"c": GREY_TXT})]],
        size=9, name="M2 legend reported")
box(slide, R0 + 2.17, LG2_Y + 0.035, 0.11, 0.11, fill=RED_BAR, shape=MSO_SHAPE.RECTANGLE, name="M2 legend core key")
textbox(slide, R0 + 2.35, LG2_Y, 2.4, 0.18,
        [[("Core OCF, ex. receivable runoff", {"c": GREY_TXT}), ("3", {"c": GREY_TXT, "sup": True})]],
        size=9, name="M2 legend core")

# brackets showing the receivable runoff gap
for k, (rep, core, runoff) in enumerate(zip(m2_rep, m2_core, ["$663M", "$489M"])):
    red_right = p2_x0 + slot2 * k + margin2 + 2 * bar_w2
    top, bot = y2(rep), y2(core)
    br = slide.shapes.add_shape(MSO_SHAPE.RIGHT_BRACE, Inches(red_right + 0.03), Inches(top),
                                Inches(0.09), Inches(bot - top))
    br.adjustments[0] = 0.04
    br.adjustments[1] = 0.5
    br.line.color.rgb = rgb(INK)
    br.line.width = Pt(0.75)
    br.fill.background()
    no_effects(br)
    br.name = f"M2 runoff bracket {k+1}"
    mid = (top + bot) / 2
    textbox(slide, red_right + 0.16, mid - 0.19, 0.95, 0.38,
            [[("Receivable runoff", {})], [(runoff, {"b": True})]], size=9.5,
            anchor=MSO_ANCHOR.MIDDLE, name=f"M2 runoff label {k+1}")

# callout box
CO_X = 11.42
co = box(slide, CO_X, 2.12, R1 - CO_X, 1.62, fill=WHITE, line=RED, dash=MSO_LINE_DASH_STYLE.ROUND_DOT,
         line_w=1.0, radius=0.08, name="M2 callout")
co.text_frame.margin_left = co.text_frame.margin_right = Inches(0.09)
co.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
fill_tf(co.text_frame, [
    [("Cash held at ~$500M to $600M only through a ", {}),
     ("shrinking lease book and new debt", {"b": True})],
    [("S&P expects core cash flow to ", {}), ("stay negative in 2026", {"b": True})],
], 10.5, space_after=6)

# ================================================================= MODULE 3
section_header(slide, L0, ROW2, L1 - L0, "Liquidity vs. Upcoming Maturities", "M3 header")

m3_cats = ["Liquidity (6/30/26)", "2028 Maturities", "2029 Maturities", "2030 Maturities"]
m3_cash = [529, None, None, None]
m3_abl = [276, None, None, None]
m3_mat = [None, 752, 1177, 1050]
C3 = dict(x=L0, y=4.55, w=4.70, h=1.80, px=0.02, pw=0.96, py=0.278, ph=0.59, vmax=1400, gap=60)


def m3_xml(rid):
    body = (f'<c:barChart><c:barDir val="col"/><c:grouping val="stacked"/><c:varyColors val="0"/>'
            + bar_ser(0, "Cash and ST investments ($M)", m3_cats, m3_cash, INK,
                      dlbls('\\$#,##0"M"', "ctr", 10, WHITE, True))
            + bar_ser(1, "ABL net availability ($M)", m3_cats, m3_abl, GREY_LT,
                      dlbls('\\$#,##0"M"', "ctr", 10, INK, True))
            + bar_ser(2, "Debt maturities ($M)", m3_cats, m3_mat, RED_BAR,
                      dlbls("", None, 10, INK, show=False))
            + f'<c:gapWidth val="{C3["gap"]}"/><c:overlap val="100"/><c:axId val="3001"/><c:axId val="3002"/></c:barChart>'
            + cat_ax(3001, 3002, line_color=GREY_MID)
            + val_ax(3002, 3001, 0, C3["vmax"]))
    return chart_space((C3["px"], C3["py"], C3["pw"], C3["ph"]), body, rid)


add_chart(slide, C3["x"], C3["y"], C3["w"], C3["h"], m3_cats,
          [("Cash and ST investments ($M)", m3_cash), ("ABL net availability ($M)", m3_abl),
           ("Debt maturities ($M)", m3_mat)], m3_xml, "M3 liquidity vs maturities chart")

p3_x0 = C3["x"] + C3["px"] * C3["w"]
p3_w = C3["pw"] * C3["w"]
p3_y0 = C3["y"] + C3["py"] * C3["h"]
p3_h = C3["ph"] * C3["h"]
p3_bot = p3_y0 + p3_h
slot3 = p3_w / 4
bar_w3 = slot3 / (1 + C3["gap"] / 100)
centers3 = [p3_x0 + slot3 * (k + 0.5) for k in range(4)]


def y3(v):
    return p3_bot - p3_h * v / C3["vmax"]


totals = [805, 752, 1177, 1050]
for k, v in enumerate(totals):
    textbox(slide, centers3[k] - 0.5, y3(v) - 0.23, 1.0, 0.2,
            [[(f"${v:,}M", {"b": True})]], size=10.5, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.BOTTOM, name=f"M3 total label {k+1}")

# dashed reference line at current liquidity
ly = y3(805)
line(slide, centers3[0] + bar_w3 / 2, ly, centers3[3] + bar_w3 / 2 + 0.05, ly, GREY_MID, 1.0,
     dash=MSO_LINE_DASH_STYLE.DASH, name="M3 liquidity reference line")
textbox(slide, centers3[3] + bar_w3 / 2 + 0.09, ly - 0.09, 0.8, 0.18,
        [[("Liquidity", {"c": GREY_TXT, "i": True}), ("4", {"c": GREY_TXT, "i": True, "sup": True})]],
        size=9, anchor=MSO_ANCHOR.MIDDLE, name="M3 liquidity line label")

# ABL springing maturity pill, tied to the 2028 bar
ABL_X, ABL_Y, ABL_W, ABL_H = 0.30, 4.55, 2.30, 0.42
abl = box(slide, ABL_X, ABL_Y, ABL_W, ABL_H, fill=WHITE, line=RED, dash=MSO_LINE_DASH_STYLE.ROUND_DOT,
          line_w=1.0, radius=0.5, name="M3 ABL springing maturity pill")
abl.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
abl.text_frame.margin_left = abl.text_frame.margin_right = Inches(0.1)
fill_tf(abl.text_frame, [[("ABL matures ", {}), ("May 2028", {"b": True, "c": RED}),
                          (", or 91 days ahead of the 2028 notes", {})]], 10, align=PP_ALIGN.CENTER)
line(slide, centers3[1], ABL_Y + ABL_H, centers3[1], y3(752) - 0.24, RED, 0.75, name="M3 ABL connector")

# coverage note in the right strip
textbox(slide, 5.02, 4.62, L1 - 5.02, 0.62,
        [[("2028 to 2030 maturities total ", {}), ("$2,979M", {"b": True, "c": RED}),
          (", ", {}), ("3.7x", {"b": True, "c": RED}), (" current liquidity", {})]],
        size=10.5, name="M3 coverage note")

LG3_X, LG3_Y = 5.05, 5.72
box(slide, LG3_X, LG3_Y + 0.035, 0.11, 0.11, fill=INK, shape=MSO_SHAPE.RECTANGLE, name="M3 legend cash key")
textbox(slide, LG3_X + 0.18, LG3_Y, 1.5, 0.18, [[("Cash and ST investments", {"c": GREY_TXT})]], size=9,
        name="M3 legend cash")
box(slide, LG3_X, LG3_Y + 0.255, 0.11, 0.11, fill=GREY_LT, shape=MSO_SHAPE.RECTANGLE, name="M3 legend ABL key")
textbox(slide, LG3_X + 0.18, LG3_Y + 0.22, 1.5, 0.18, [[("ABL net availability", {"c": GREY_TXT})]], size=9,
        name="M3 legend ABL")
box(slide, LG3_X, LG3_Y + 0.475, 0.11, 0.11, fill=RED_BAR, shape=MSO_SHAPE.RECTANGLE, name="M3 legend maturities key")
textbox(slide, LG3_X + 0.18, LG3_Y + 0.44, 1.5, 0.18, [[("Debt maturities", {"c": GREY_TXT}), ("5", {"c": GREY_TXT, "sup": True})]], size=9,
        name="M3 legend maturities")

# ================================================================= MODULE 4
section_header(slide, R0, ROW2, R1 - R0, "Key Business and Covenant Issues", "M4 header")

items = [
    [[("> ", {}), ("Secular decline: ", {"b": True}),
      ("print is ~90% of revenue vs. ~54% at Canon and ~30% at HP; pro forma revenue ", {}),
      ("(6.5%)", {"b": True, "c": RED}), (" in Q2'26", {})]],
    [[("> ", {}), ("Interest burden: ", {"b": True}),
      ("~$375M of 2026E non-financing interest vs. $614M of LTM Adj. EBITDA ex-tariff", {}),
      ("2", {"sup": True}), (" (", {}), ("~1.6x", {"b": True, "c": RED}), (" coverage)", {})]],
    [[("> ", {}), ("Covenants: ", {"b": True}), ("springing ", {}), ("1.0x FCCR", {"b": True}),
      (" if ABL availability falls below the greater of $31.9M or 10% of the borrowing base", {})],
     [("$2B secured cap", {"b": True}), (" and ", {}), ("$370M non-guarantor cap", {"b": True}),
      (" limit new money; rated ", {}), ("SD", {"b": True, "c": RED}), (" in June 2026", {})]],
]
NUM_D = 0.24
y = ROW2 + 0.46
heights = [0.40, 0.40, 0.62]
for k, paras in enumerate(items):
    circ = box(slide, R0, y + 0.01, NUM_D, NUM_D, fill=INK, shape=MSO_SHAPE.OVAL, name=f"M4 number {k+1}")
    circ.text_frame.margin_left = circ.text_frame.margin_right = 0
    circ.text_frame.margin_top = circ.text_frame.margin_bottom = 0
    circ.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_tf(circ.text_frame, [[(str(k + 1), {"b": True, "c": WHITE})]], 11, align=PP_ALIGN.CENTER)
    tb = textbox(slide, R0 + 0.36, y, R1 - R0 - 0.36, heights[k], paras, size=11,
                 name=f"M4 item {k+1}", space_after=2)
    if len(paras) > 1:
        sub = tb.text_frame.paragraphs[1]
        sub.level = 0
        pPr = sub._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.14)))
        first = sub.runs[0]
        r = copy.deepcopy(first._r)
        r.find(qn("a:t")).text = "> "
        r.find(qn("a:rPr")).set("b", "0")
        sub._p.insert(list(sub._p).index(first._r), r)
    y += heights[k] + 0.10

# ================================================================= CEO callout
ceo = box(slide, L0, 6.43, R1 - L0, 0.36, fill=WHITE, line=INK, dash=MSO_LINE_DASH_STYLE.DASH,
          line_w=0.75, radius=0.2, name="CEO three hurdles callout")
ceo.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
fill_tf(ceo.text_frame, [[
    ("CEO Louie Pastor frames the ", {"i": True}),
    ("2028", {"b": True, "c": RED}), (", ", {"i": True}),
    ("2029", {"b": True, "c": RED}), (" and ", {"i": True}),
    ("2030", {"b": True, "c": RED}),
    (" maturities as “three hurdles”: if an initiative does not advance one of them, Xerox does not pursue it",
     {"i": True}),
    ("6", {"i": True, "sup": True}),
]], 11.5, align=PP_ALIGN.CENTER)

prs.save(OUT)
print("saved", OUT)
