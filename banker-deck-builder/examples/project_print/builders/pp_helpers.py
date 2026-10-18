"""Add the Situation Overview slide (slide 4) to Project Print deck v2.

Usage: python3 build_situation.py <input.pptx> <output.pptx>
"""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt



# ---------------------------------------------------------------- palette
INK = "1A1A24"
RED = "D92332"      # lines, markers, pills, red text (matches slide 3)
RED_BAR = "C00000"  # bar fills (matches the Xerox bar on slide 3)
GREY_TXT = "595959"
GREY_MID = "7F8590"
GREY_LT = "BFC3C9"
WHITE = "FFFFFF"
FONT = "Garamond"


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- text helpers
def style_run(run, size, bold=False, italic=False, color=INK, sup=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    if sup:
        run._r.get_or_add_rPr().set("baseline", "30000")


def fill_tf(tf, paras, size, align=PP_ALIGN.LEFT, space_after=0, line_spacing=None):
    """paras: list of paragraphs; each paragraph is a list of (text, opts) runs."""
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, o in runs:
            r = p.add_run()
            r.text = text
            style_run(r, o.get("size", size), o.get("b", False), o.get("i", False),
                      o.get("c", INK), o.get("sup", False))


def textbox(slide, x, y, w, h, paras, size=11, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, name=None, space_after=0, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    fill_tf(tf, paras, size, align, space_after)
    if name:
        tb.name = name
    return tb


def no_effects(shape):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    spPr.append(etree.SubElement(spPr, qn("a:effectLst")))


def box(slide, x, y, w, h, fill=WHITE, line=None, dash=None, line_w=0.75,
        radius=None, name=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = dash
    no_effects(s)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.word_wrap = True
    if name:
        s.name = name
    return s


def line(slide, x1, y1, x2, y2, color=INK, width=1.0, dash=None, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if name:
        c.name = name
    return c


def section_header(slide, x, y, w, text, name):
    textbox(slide, x, y, w, 0.24, [[(text, {"b": True})]], size=14, name=name)
    line(slide, x, y + 0.29, x + w, y + 0.29, INK, 1.0, name=name + " rule")


# ---------------------------------------------------------------- chart XML
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def col(i):
    return chr(ord("B") + i)


def txpr(size, color, bold=False):
    b = ' b="1"' if bold else ' b="0"'
    return (f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="{int(size*100)}"{b}>'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="{FONT}"/></a:defRPr></a:pPr><a:endParaRPr lang="en-US"/></a:p></c:txPr>')


def str_cache(vals):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals))
    return f'<c:strCache><c:ptCount val="{len(vals)}"/>{pts}</c:strCache>'


def num_cache(vals, fmt="General"):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(vals)
                  if v is not None)
    return f'<c:numCache><c:formatCode>{fmt}</c:formatCode><c:ptCount val="{len(vals)}"/>{pts}</c:numCache>'


def dlbls(fmt, pos, size, color, bold=True, show=True, deleted_idx=()):
    if not show:
        return '<c:dLbls><c:delete val="1"/></c:dLbls>'
    dels = "".join(f'<c:dLbl><c:idx val="{i}"/><c:delete val="1"/></c:dLbl>' for i in deleted_idx)
    pos_xml = f'<c:dLblPos val="{pos}"/>' if pos else ""
    fmt = fmt.replace('&', '&amp;').replace('"', '&quot;')
    return (f'<c:dLbls>{dels}<c:numFmt formatCode="{fmt}" sourceLinked="0"/>'
            f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>{txpr(size, color, bold)}'
            f'{pos_xml}<c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/>'
            f'<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>')


def bar_ser(i, name, cats, vals, fill, labels, dpt_fills=None):
    n = len(cats)
    dpts = ""
    for idx, f in (dpt_fills or {}).items():
        dpts += (f'<c:dPt><c:idx val="{idx}"/><c:invertIfNegative val="0"/><c:bubble3D val="0"/>'
                 f'<c:spPr><a:solidFill><a:srgbClr val="{f}"/></a:solidFill></c:spPr></c:dPt>')
    return (f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/>'
            f'<c:tx><c:strRef><c:f>Sheet1!${col(i)}$1</c:f>{str_cache([name])}</c:strRef></c:tx>'
            f'<c:spPr><a:solidFill><a:srgbClr val="{fill}"/></a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>'
            f'<c:invertIfNegative val="0"/>{dpts}{labels}'
            f'<c:cat><c:strRef><c:f>Sheet1!$A$2:$A${n+1}</c:f>{str_cache(cats)}</c:strRef></c:cat>'
            f'<c:val><c:numRef><c:f>Sheet1!${col(i)}$2:${col(i)}${n+1}</c:f>{num_cache(vals)}</c:numRef></c:val>'
            f'</c:ser>')


def line_ser(i, name, cats, vals, color, labels, marker="diamond", msize=6, mfill=WHITE,
             show_line=True, mline_w=1.0):
    n = len(cats)
    ln = (f'<a:ln w="19050" cap="rnd"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:ln>'
          if show_line else '<a:ln><a:noFill/></a:ln>')
    mfill_xml = (f'<a:solidFill><a:srgbClr val="{mfill}"/></a:solidFill>' if mfill
                 else '<a:noFill/>')
    return (f'<c:ser><c:idx val="{i}"/><c:order val="{i}"/>'
            f'<c:tx><c:strRef><c:f>Sheet1!${col(i)}$1</c:f>{str_cache([name])}</c:strRef></c:tx>'
            f'<c:spPr>{ln}</c:spPr>'
            f'<c:marker><c:symbol val="{marker}"/><c:size val="{msize}"/>'
            f'<c:spPr>{mfill_xml}<a:ln w="{int(mline_w*12700)}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:ln></c:spPr></c:marker>'
            f'{labels}'
            f'<c:cat><c:strRef><c:f>Sheet1!$A$2:$A${n+1}</c:f>{str_cache(cats)}</c:strRef></c:cat>'
            f'<c:val><c:numRef><c:f>Sheet1!${col(i)}$2:${col(i)}${n+1}</c:f>{num_cache(vals)}</c:numRef></c:val>'
            f'<c:smooth val="0"/></c:ser>')


def cat_ax(ax_id, cross_id, delete=False, lbl_pos="nextTo", line_color=GREY_LT, size=10):
    if delete:
        return (f'<c:catAx><c:axId val="{ax_id}"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
                f'<c:delete val="1"/><c:axPos val="b"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
                f'<c:tickLblPos val="nextTo"/><c:crossAx val="{cross_id}"/><c:crosses val="autoZero"/>'
                f'<c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>')
    return (f'<c:catAx><c:axId val="{ax_id}"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
            f'<c:delete val="0"/><c:axPos val="b"/><c:numFmt formatCode="General" sourceLinked="0"/>'
            f'<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="{lbl_pos}"/>'
            f'<c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="{line_color}"/></a:solidFill></a:ln></c:spPr>'
            f'{txpr(size, GREY_TXT)}<c:crossAx val="{cross_id}"/><c:crosses val="autoZero"/>'
            f'<c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>')


def val_ax(ax_id, cross_id, vmin, vmax, pos="l", crosses="autoZero", hidden_visible=False):
    if hidden_visible:
        return (f'<c:valAx><c:axId val="{ax_id}"/><c:scaling><c:orientation val="minMax"/>'
                f'<c:max val="{vmax}"/><c:min val="{vmin}"/></c:scaling><c:delete val="0"/>'
                f'<c:axPos val="{pos}"/><c:numFmt formatCode="General" sourceLinked="0"/>'
                f'<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="none"/>'
                f'<c:spPr><a:ln><a:noFill/></a:ln></c:spPr>'
                f'<c:crossAx val="{cross_id}"/><c:crosses val="{crosses}"/><c:crossBetween val="between"/></c:valAx>')
    return (f'<c:valAx><c:axId val="{ax_id}"/><c:scaling><c:orientation val="minMax"/>'
            f'<c:max val="{vmax}"/><c:min val="{vmin}"/></c:scaling><c:delete val="1"/>'
            f'<c:axPos val="{pos}"/><c:numFmt formatCode="General" sourceLinked="0"/>'
            f'<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
            f'<c:crossAx val="{cross_id}"/><c:crosses val="{crosses}"/><c:crossBetween val="between"/></c:valAx>')


def chart_space(plot_inner, plot_body, rid):
    x, y, w, h = plot_inner
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<c:chartSpace xmlns:c="{C_NS}" xmlns:a="{A_NS}" xmlns:r="{R_NS}">'
            f'<c:date1904 val="0"/><c:lang val="en-US"/><c:roundedCorners val="0"/>'
            f'<c:chart><c:autoTitleDeleted val="1"/><c:plotArea>'
            f'<c:layout><c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
            f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{h}"/></c:manualLayout></c:layout>'
            f'{plot_body}<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr></c:plotArea>'
            f'<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/></c:chart>'
            f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
            f'{txpr(10, GREY_TXT)}'
            f'<c:externalData r:id="{rid}"><c:autoUpdate val="0"/></c:externalData></c:chartSpace>')


def add_chart(slide, x, y, w, h, cats, series, xml_builder, name):
    """Create a native chart (embedded workbook holds every series), then replace its XML."""
    cd = CategoryChartData()
    cd.categories = cats
    for s_name, vals in series:
        cd.add_series(s_name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
    gf.name = name
    part = gf.chart_part
    rid = part._element.find(qn("c:externalData")).get(qn("r:id"))
    part._element = parse_xml(xml_builder(rid).encode("utf-8"))
    return gf


