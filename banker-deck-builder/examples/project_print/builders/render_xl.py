"""Render Excel ranges from the recalculated model as PNG images (for appendix 'Excel output' pages).

Usage: python3 render_xl.py <job> [<job> ...]   (job names from JOBS; 'all' renders every job)
Values come from model/v16_verify.xlsx (LibreOffice-recalculated; FactSet values injected); the DCF sensitivity
data tables take their values from model/Xerox_RX_Model_v16.xlsx.
"""
import os
import shutil
import subprocess
import sys

from openpyxl import load_workbook
from openpyxl.worksheet.page import PageMargins

VER = "model/v16_verify.xlsx"
REAL = "model/Xerox_RX_Model_v16.xlsx"
OUTDIR = "xl_img"

JOBS = {
    "rev": ("$XRX Revenue Build", "C8:N35"),
    "opa": ("$XRX DCF Operating Assumptions", "C8:N37"),
    "dcf_fcf": ("$XRX DCF", "C5:N58"),
    "dcf_grid_x": ("$XRX DCF", "C68:K86"),
    "dcf_top": ("$XRX DCF", "C5:N17"),
    "dcf_wacc": ("$XRX DCF", "C19:H41"),
    "dcf_ev": ("$XRX DCF", "C44:N58"),
    "liq_a": ("$XRX Liquidation Analysis", "C8:J35"),
    "liq_b": ("$XRX Liquidation Analysis", "C38:J57"),
    "dcf_grid_g": ("$XRX DCF", "C94:K112"),
    "comps": ("$XRX Comparable Companies", "C5:X29"),
    "prec": ("Precedent Transactions", "C7:M33"),
    "liq": ("$XRX Liquidation Analysis", "C8:J57"),
    "wf": ("$XRX Recovery Waterfall", "C63:J84"),
    "wf_dv": ("$XRX Recovery Waterfall", "C22:I30"),
    "lq_fcf": ("$XRX Liquidity Analysis", "C57:I75"),
    "lq_cash": ("$XRX Liquidity Analysis", "C121:I140"),
    "tx_su": ("$XRX Transaction Analysis", "C38:E49"),
    "tx_pf": ("$XRX Transaction Analysis", "C52:F74"),
    "tx_eq": ("$XRX Transaction Analysis", "C77:G98"),
    "tx_b": ("$XRX Transaction Analysis", "C129:I154"),
}


def render(job, dpi=250):
    sheet, rng = JOBS[job]
    wb = load_workbook(VER, data_only=True)
    if sheet == "$XRX DCF":
        real = load_workbook(REAL)[sheet]
        d = wb[sheet]
        for r0 in (72, 82, 98, 108):
            for r in range(r0, r0 + 5):
                for c in "GHIJK":
                    v = real[f"{c}{r}"].value
                    if isinstance(v, (int, float)):
                        d[f"{c}{r}"].value = v
        # data-table anchor cells (top-left of each grid) hold the TABLE formula: compute them directly
        uf = [d[f"{c}13"].value for c in "JKLMN"]
        st = [d[f"{c}16"].value for c in "JKLMN"]
        per = [d[f"{c}15"].value for c in "JKLMN"]
        eb, w0, x0, g0, debt = d["G45"].value, d["G41"].value, d["G46"].value, d["N45"].value, d["G34"].value
        pv = lambda w: sum(u * s_ / (1 + w) ** p for u, s_, p in zip(uf, st, per))
        ev_x = pv(w0 + 0.02) + eb * (x0 - 1.0) / (1 + w0 + 0.02) ** (per[-1] + 0.5)
        ev_g = pv(w0 + 0.02) + uf[-1] * (1 + g0 - 0.01) / (w0 + 0.02 - g0 + 0.01) / (1 + w0 + 0.02) ** per[-1]
        d["G72"].value, d["G82"].value = ev_x, ev_x - debt
        d["G98"].value, d["G108"].value = ev_g, ev_g - debt
    for n in wb.sheetnames:
        if n != sheet:
            del wb[n]
    ws = wb[sheet]
    ws.print_area = rng
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.gridLines = False
    ws.print_options.headings = False
    ws.page_margins = PageMargins(left=0.1, right=0.1, top=0.1, bottom=0.1, header=0, footer=0)
    ws.oddHeader.center.text = ws.oddFooter.center.text = None
    ws.oddHeader.left.text = ws.oddHeader.right.text = ws.oddFooter.left.text = ws.oddFooter.right.text = None
    tmp = os.path.join(OUTDIR, f"_{job}")
    os.makedirs(tmp, exist_ok=True)
    x = os.path.join(tmp, f"{job}.xlsx")
    wb.save(x)
    subprocess.run(["timeout", "120", "soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, x],
                   capture_output=True)
    pdf = os.path.join(tmp, f"{job}.pdf")
    base = os.path.join(tmp, "p")
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", "-f", "1", "-l", "1", "-singlefile", pdf, base], check=True)
    out = os.path.join(OUTDIR, f"{job}.png")
    subprocess.run(["convert", base + ".png", "-trim", "+repage", "-bordercolor", "white", "-border", "6", out],
                   check=True)
    shutil.rmtree(tmp, ignore_errors=True)
    print("rendered", job, out)


if __name__ == "__main__":
    os.makedirs(OUTDIR, exist_ok=True)
    jobs = list(JOBS) if sys.argv[1:] == ["all"] else sys.argv[1:]
    for j in jobs:
        render(j)
