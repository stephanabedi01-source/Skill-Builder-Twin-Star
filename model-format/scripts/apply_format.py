#!/usr/bin/env python3
"""
apply_format.py - apply the house style from a (reviewed) JSON map.

Consumes the corrected audit map and applies every visual convention through
openpyxl, pulling EVERY hex code / number-format / border from style_constants.py
(never from memory). Guarantees:

  * NEVER changes a cell's value or formula -- only .font/.fill/.number_format/
    .alignment/.border, plus sheet-level gridlines/freeze/tab-color/page-setup.
  * NEVER inserts, deletes, or reorders rows/columns/tabs.
  * Skips tabs whose map says format_body=false (working/source/divider/hidden) --
    their bodies are left exactly as-is (a working tab keeps its gridlines).
  * Idempotent: running twice produces an identical file. Borders are merged onto
    existing borders, and every other attribute is set to a fixed value, so a
    second pass changes nothing.

Font color is derived live from each cell's provenance (hardcode->blue,
cross-sheet->green, formula/text->gray), so the map only needs to carry the
semantic judgment (tab type, row class, tier, input blocks) -- not a color per
cell.

Usage:
    python apply_format.py TARGET.xlsx --map audit_map.json [-o OUT.xlsx]
"""

import argparse
import json
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.properties import PageSetupProperties

import style_constants as sc
from audit_workbook import provenance

TOTAL_CLASSES = {"subtotal", "major_total", "headline_total"}
TIER_FILL = {"subtotal": sc.FILL_TIER1, "major_total": sc.FILL_TIER2, "headline_total": sc.FILL_TIER3}

# class -> (font color, bold, italic) when the class overrides provenance entirely
FORCED_FONT = {
    "banner":           (sc.WHITE, True, False),
    "subbanner":        (sc.WHITE, True, False),
    "subsection_label": (sc.SUBLABEL_BLUE, True, False),
    "unit_note":        (sc.NOTE_PURPLE, False, True),
    "check":            (sc.FLAG_RED, False, True),
    "percent":          (sc.SLATE_GRAY, False, True),
    "header_fy":        (sc.WHITE, True, False),
}
NUMFMT_BY_CLASS = {
    "data_first":      sc.NF_DOLLAR_FIRST,
    "subtotal":        sc.NF_TOTAL,
    "major_total":     sc.NF_TOTAL,
    "headline_total":  sc.NF_TOTAL,
    "percent":         sc.NF_PERCENT_MARGIN,
    "check":           sc.NF_CHECK,
    "header_date":     sc.NF_DATE_MONTH,
}
INPUT_NUMFMT = {
    "percent_input": sc.NF_PERCENT_INPUT,
    "capex_input":   sc.NF_CAPEX_INPUT,
    "rate_spread":   sc.NF_RATE_SPREAD,
}
CENTER_CLASSES = {"header_date", "header_fy", "actual_forecast", "toggle"}
WARNING_RE = re.compile(
    r"CONFIDENTIAL|DO NOT (DISTRIBUTE|COPY|FORWARD)|PRIVILEGED|"
    r"SUBJECT TO (MATERIAL )?REVISION|PROPRIETARY|FOR DISCUSSION", re.I)


def _idx(letter):
    return column_index_from_string(letter)


def _numeric_or_blank(cell):
    """True if it's safe to stamp a number format here (number/formula/empty)."""
    return cell.data_type in ("n", "f") or cell.value is None


# ---------------------------------------------------------------------------
# Border merging (idempotent)
# ---------------------------------------------------------------------------
def _add_total_border(ws, row, c0, c1):
    g = sc._side(sc.BORDER_GRAY)
    for col in range(c0, c1 + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = sc.merge_border(cell.border, top=g, bottom=g)
        if row > 1:
            above = ws.cell(row=row - 1, column=col)
            above.border = sc.merge_border(above.border, bottom=g)


def _add_input_box(ws, r0, r1, c0, c1):
    b = sc._side(sc.BORDER_BLUE)
    for row in range(r0, r1 + 1):
        for col in range(c0, c1 + 1):
            cell = ws.cell(row=row, column=col)
            sides = {}
            if row == r0:
                sides["top"] = b
            if row == r1:
                sides["bottom"] = b
            if col == c0:
                sides["left"] = b
            if col == c1:
                sides["right"] = b
            if sides:
                cell.border = sc.merge_border(cell.border, **sides)


# ---------------------------------------------------------------------------
# Page / sheet chrome
# ---------------------------------------------------------------------------
def _setup_presentation_page(ws):
    ws.sheet_view.showGridLines = False
    ps = ws.page_setup
    if ps.scale or ps.fitToWidth or ps.orientation or ps.paperSize:
        # The sheet already carries deliberate print setup (e.g. an explicit
        # per-tab scale). Imposing fit-to-page or paper defaults on top of it is
        # noise that diverges from the source -- leave print setup alone.
        return
    ps.orientation = sc.PAGE_ORIENTATION
    ps.paperSize = 1  # Letter
    ps.fitToWidth = 1  # "label columns plus period fit one page wide"
    ps.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    pm = ws.page_margins
    pm.left = pm.right = sc.MARGIN_LR
    pm.top = pm.bottom = sc.MARGIN_TB


def _set_tab_color(ws, hex6):
    if hex6:
        ws.sheet_properties.tabColor = sc._argb(hex6)


# ---------------------------------------------------------------------------
# Per-sheet application
# ---------------------------------------------------------------------------
def _format_sheet(ws, sheet):
    label_cols = sheet.get("label_cols") or []
    marker = _idx(sheet["marker_col"]) if sheet.get("marker_col") else None
    c_start = _idx(sheet["data_col_start"])
    c_end = _idx(sheet["data_col_end"])
    label_idx_set = {_idx(x) for x in label_cols}
    first_label = min(label_idx_set, default=c_start)
    span_lo, span_hi = first_label, c_end  # banners / total fills span label..last data col
    balance = sheet.get("numfmt_family") == "balance"
    fam = sc.HEADER_FILLS.get(sheet.get("header_family", "core"), sc.HEADER_FILLS["core"])
    fy_fill, date_fill = fam
    font_name = sheet.get("font_family", sc.FONT_BODY)
    is_cover = sheet.get("tab_type") == "cover"
    # Output / presentation pages are intentionally all-gray: provenance coloring
    # (blue inputs / green links) is dropped on them. See references/tab-patterns.md.
    gray_only = sheet.get("tab_type") == "output"

    rowmap = {r["row"]: r for r in sheet.get("rows", [])}

    # Input-block cell membership.
    input_cells = set()
    input_numfmt = {}
    for blk in sheet.get("input_blocks", []):
        r0, r1 = blk["min_row"], blk["max_row"]
        b0, b1 = _idx(blk["min_col"]), _idx(blk["max_col"])
        nf = INPUT_NUMFMT.get(blk.get("numfmt"))
        for rr in range(r0, r1 + 1):
            for cc in range(b0, b1 + 1):
                input_cells.add((rr, cc))
                if nf:
                    input_numfmt[(rr, cc)] = nf

    # Full used range -- impose the base font everywhere, class styling where mapped.
    min_col, min_row, max_col, max_row = range_boundaries(ws.calculate_dimension())
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            col = cell.column
            r = cell.row
            info = rowmap.get(r)
            cls = info["class"] if info else None
            is_marker = marker is not None and col == marker
            is_data = c_start <= col <= c_end
            in_span = span_lo <= col <= span_hi
            in_input = (r, col) in input_cells

            # ---- FONT ----
            color, bold, italic, underline = sc.SLATE_GRAY, False, False, None
            forced = None
            if cls in FORCED_FONT and (in_span or is_data or col in label_idx_set):
                color, bold, italic = FORCED_FONT[cls]
                forced = cls
            elif cls == "section_underline":
                # bold + underlined gray header (BS "Assets" / "Liabilities & Equity")
                color, bold, underline, forced = sc.SLATE_GRAY, True, "single", cls
            elif cls in TOTAL_CLASSES:
                bold = True
            elif cls == "header_date":
                bold = True

            if not forced and not gray_only:
                # provenance drives the color of any cell carrying content in the data area,
                # and of input cells (so green links inside a toggle stay green).
                if (is_data or in_input) and cell.value is not None:
                    p = provenance(cell)
                    color = (sc.LINK_GREEN if p == "cross_sheet"
                             else sc.INPUT_BLUE if p == "hardcode"
                             else sc.SLATE_GRAY)
            if in_input and not gray_only and not is_marker:
                # A triple-marked input is always provenance-colored and never bold/italic,
                # even if its row was classed as percent/total -- input membership wins.
                p = provenance(cell)
                color = (sc.LINK_GREEN if p == "cross_sheet"
                         else sc.INPUT_BLUE if p == "hardcode" else sc.SLATE_GRAY)
                bold = italic = False
                underline = None
            if is_marker and cell.value is not None:
                color, bold = sc.FLAG_RED, True  # red "X" navigation markers
            if is_cover and isinstance(cell.value, str) and WARNING_RE.search(cell.value):
                color, bold = sc.FLAG_RED, True  # confidentiality warnings on the cover

            cell.font = sc.font(color, bold=bold, italic=italic, underline=underline, name=font_name)

            # ---- FILL ---- (only where a class/input dictates; never clear others)
            if cls == "banner" and in_span:
                cell.fill = sc.fill(sc.FILL_BANNER)
            elif cls == "subbanner" and in_span:
                cell.fill = sc.fill(sc.FILL_SUBBANNER)
            elif cls in TIER_FILL and in_span:
                cell.fill = sc.fill(TIER_FILL[cls])
            elif cls == "header_fy" and is_data:
                cell.fill = sc.fill(fy_fill)
            elif cls == "header_date" and is_data:
                cell.fill = sc.fill(date_fill)
            elif in_input:
                cell.fill = sc.fill(sc.FILL_INPUT)

            # ---- NUMBER FORMAT ---- (data columns; never on text cells)
            if is_data and _numeric_or_blank(cell):
                if in_input and (r, col) in input_numfmt:
                    cell.number_format = input_numfmt[(r, col)]
                elif cls == "data":
                    cell.number_format = sc.NF_NUMBER_BS if balance else sc.NF_NUMBER
                elif cls == "header_date":
                    # Stamp the date format on real dates / =EOMONTH formulas, and on
                    # bare date *serials* (>=20000 ~ year 1954, e.g. a stripped model
                    # whose dates lost their format) -- but never on a plain year
                    # integer like 2019 (< 20000), which would become a 1905 serial.
                    if (cell.is_date or cell.data_type == "f"
                            or (cell.data_type == "n" and isinstance(cell.value, (int, float))
                                and not isinstance(cell.value, bool) and cell.value >= 20000)):
                        cell.number_format = sc.NF_DATE_MONTH
                elif cls in NUMFMT_BY_CLASS:
                    cell.number_format = NUMFMT_BY_CLASS[cls]

            # ---- ALIGNMENT ---- (centered header classes; map-driven indent
            #      hierarchy on label cells; everything else preserved)
            if cls in CENTER_CLASSES:
                cell.alignment = sc.CENTER
            elif info and info.get("indent") and col in label_idx_set:
                cell.alignment = sc.left_indent(info["indent"])

    # ---- BORDERS ----
    for r, info in rowmap.items():
        if info["class"] in TOTAL_CLASSES:
            _add_total_border(ws, r, span_lo, span_hi)
    for blk in sheet.get("input_blocks", []):
        _add_input_box(ws, blk["min_row"], blk["max_row"], _idx(blk["min_col"]), _idx(blk["max_col"]))


def _apply_geometry(ws, sheet):
    """Mode-B geometry: chassis column widths / row heights. Only runs when the
    reviewed map opts in ("auto_widths": true and/or explicit "column_widths" /
    "row_heights"), so replicate-mode and legacy maps are untouched."""
    widths = {}
    if sheet.get("auto_widths"):
        label_idx = {_idx(x) for x in (sheet.get("label_cols") or [])}
        marker = _idx(sheet["marker_col"]) if sheet.get("marker_col") else None
        c_start = _idx(sheet["data_col_start"])
        c_end = _idx(sheet["data_col_end"])
        first_label = min(label_idx, default=c_start)
        for col in range(1, c_end + 1):
            if col in label_idx:
                widths[col] = sc.WIDTH_LABEL
            elif col == marker or col < first_label:
                widths[col] = sc.WIDTH_MARGIN
            elif c_start <= col <= c_end:
                widths[col] = sc.WIDTH_DATA
    for letter, w in (sheet.get("column_widths") or {}).items():
        widths[_idx(letter)] = float(w)
    for col, w in widths.items():
        # setting width is enough -- openpyxl derives customWidth from it
        ws.column_dimensions[get_column_letter(col)].width = float(w)
    for r, h in (sheet.get("row_heights") or {}).items():
        ws.row_dimensions[int(r)].height = float(h)


def apply_map(wb_path, map_path, out_path):
    with open(map_path) as f:
        m = json.load(f)
    wb = load_workbook(wb_path, data_only=False)
    by_name = {s["name"]: s for s in m["sheets"]}

    formatted, skipped = [], []
    for ws in wb.worksheets:
        sheet = by_name.get(ws.title)
        if sheet is None:
            skipped.append((ws.title, "not in map"))
            continue
        # Tab color is applied even to unformatted (working/source) tabs.
        _set_tab_color(ws, sheet.get("tab_color"))
        if not sheet.get("format_body"):
            skipped.append((ws.title, sheet.get("tab_type", "?")))
            continue
        if not sheet.get("data_col_start"):
            # Hand-edited map turned on format_body without geometry -- skip safely.
            skipped.append((ws.title, "format_body set but geometry missing"))
            continue
        if sheet.get("gridlines_off"):
            _setup_presentation_page(ws)
        fr = sheet.get("freeze")
        if fr and (ws.freeze_panes in (None, "A1") or sheet.get("freeze_force")):
            # Never move a freeze the source already has -- a deliberate freeze
            # (e.g. 76 frozen history columns on a liquidity tab) encodes intent
            # a header heuristic cannot reconstruct. Override only with
            # "freeze_force": true in the reviewed map.
            ws.freeze_panes = fr
        _format_sheet(ws, sheet)
        _apply_geometry(ws, sheet)
        formatted.append(ws.title)

    wb.save(out_path)
    return formatted, skipped


def main():
    ap = argparse.ArgumentParser(description="Apply the house model-format style from a reviewed map.")
    ap.add_argument("workbook")
    ap.add_argument("--map", required=True, help="reviewed audit_map.json")
    ap.add_argument("-o", "--output", help="output path (default: overwrite input in place)")
    args = ap.parse_args()
    out = args.output or args.workbook

    formatted, skipped = apply_map(args.workbook, args.map, out)
    print(f"Formatted {len(formatted)} tab(s): {', '.join(formatted) or '(none)'}")
    print(f"Skipped {len(skipped)} tab(s):")
    for name, why in skipped:
        print(f"   - {name}  ({why})")
    print(f"Saved -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
