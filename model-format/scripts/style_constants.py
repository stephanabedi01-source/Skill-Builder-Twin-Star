#!/usr/bin/env python3
"""
style_constants.py - SINGLE SOURCE OF TRUTH for the house model-format style.

Every hex code, number-format string, font spec, fill, and border weight in the
design language lives here as a named constant, plus small openpyxl factory
functions that build style objects from those constants. Nothing downstream
(audit_workbook.py, apply_format.py) should ever hardcode a color or a format
string -- it imports from this module. That is the mechanism that guarantees
format strings and hex codes are applied exactly, never reproduced from memory.

All colors are stored as 6-digit RGB hex (no '#', no alpha) exactly as written in
the design language. _argb() converts to the 8-digit ARGB openpyxl wants. We
always write *resolved* RGB hex (never theme indices or tints), so the style is
reproduced faithfully on any target workbook regardless of its theme.
"""

from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

# ===========================================================================
# 1. FONT COLORS  ==  DATA PROVENANCE
#    (font color encodes where a value came from; this is the most distinctive
#     trait of the style -- body text is slate gray, never pure black)
# ===========================================================================
SLATE_GRAY    = "525766"  # all formulas/calculations and labels (default body text)
INPUT_BLUE    = "0000FF"  # hardcoded inputs & pasted/loaded historical data rows
LINK_GREEN    = "00B050"  # links pulling from another worksheet (e.g. =Assumptions!$D$2)
FLAG_RED      = "FF0000"  # check rows, error flags, bold "X" nav markers, to-dos, confidentiality
NOTE_PURPLE   = "666699"  # unit notes ("#'s in thousands"), italic, once per tab
SUBLABEL_BLUE = "0067A5"  # sub-section labels (bold), e.g. "Cash Flows from Operating Activities:"
WHITE         = "FFFFFF"  # text inside dark banner fills (bold)

# ===========================================================================
# 2. FILLS  ==  3-TIER TOTAL HIERARCHY + SPECIAL FILLS
#    (all total tiers use bold text; identical on every tab)
# ===========================================================================
FILL_TIER1     = "F2F2F2"  # Tier 1: subtotals (white tint -0.05)
FILL_TIER2     = "DCE8F4"  # Tier 2: major totals -- accent #508BC9 at tint 0.8 (NOT DBE8F4; that was a one-digit-off extraction)
FILL_TIER3     = "D3EFEF"  # Tier 3: headline totals -- accent #24B1B1 at tint 0.8 (NOT CDF5F5)
FILL_INPUT     = "FFFFCC"  # input cells -- ALWAYS triple-marked (this fill + blue font + blue box)
FILL_BANNER    = "002855"  # section banner rows (dark navy), white bold text, spans label->last data col
FILL_SUBBANNER = "9FC3DA"  # secondary/sub-group banners (steel blue), white bold text

# ===========================================================================
# 3. BORDERS
# ===========================================================================
BORDER_GRAY = "BFBFBF"  # thin gray rule above & below total rows
BORDER_BLUE = "0000FF"  # thin blue box around an input block (one box per rectangle)

# ===========================================================================
# 4. HEADER FILL PAIRS BY TAB FAMILY  ->  (FY-band fill, date-row fill)
#    RULE: the date/period band under a header band is the 0.6 TINT of that
#    band's own hue -- it is not one global color. The pairs below are the
#    common families, with the date-row value precomputed via tint(hue, 0.6).
#    A model with per-block scenario bands gets per-block tints (use tint()).
# ===========================================================================
def tint(hex6, t):
    """Excel-style positive tint: each channel c -> c + (255 - c) * t."""
    h = str(hex6).lstrip("#")
    return "".join(f"{round(int(h[i:i+2], 16) + (255 - int(h[i:i+2], 16)) * t):02X}"
                   for i in (0, 2, 4))


HEADER_FILLS = {
    "core":        ("525766", "BABCC2"),  # Core statements (IS, BS, CF); BABCC2 = tint(525766, 0.6)
    "liquidity":   ("24B1B1", "A7E0E0"),  # Liquidity / borrowing base;   A7E0E0 = tint(24B1B1, 0.6)
    "ancillary":   ("BCBFC6", "E4E5E8"),  # Ancillary forecasts;          E4E5E8 = tint(BCBFC6, 0.6)
    "assumptions": ("6C1E36", "C4A5AF"),  # Assumptions "Live Case";      C4A5AF = tint(6C1E36, 0.6)
}

# ===========================================================================
# 5. TAB COLOR FAMILIES PER WORKBOOK SECTION
#    (cover / assumptions / divider tabs get NO tab color -> omitted here)
# ===========================================================================
TAB_COLORS = {
    "core_statement": "0067A5",  # IS / BS / CF
    "ancillary":      "BCBFC6",  # Sales / COGS / BS forecasts
    "carveout":       "24B1B1",  # carveout / special analyses
    "output":         "FFFF00",  # presentation output tabs
    "source":         "00B0F0",  # raw source-data tabs
    "working":        "00B050",  # working / scratch tabs
}

# ===========================================================================
# 6. NUMBER FORMATS  (EXACT STRINGS -- copy character-for-character)
#    Conventions baked in: negatives in parentheses; zeros render as "-";
#    "$" only on the first row of a block and on totals.
# ===========================================================================
NF_DOLLAR_FIRST   = "$#,##0_);($#,##0);-;"          # first $ row of a block (Gross Sales, Cash, Net Income)
NF_NUMBER         = "#,##0_);(#,##0);-;"            # subsequent rows in a block
NF_NUMBER_BS      = "#,##0_);(#,##0);-_);"          # Balance Sheet / Cash Flow variant of subsequent rows
NF_TOTAL          = "$#,##0_);($#,##0);$-_);"       # total rows (all three tiers)
NF_PERCENT_MARGIN = "0.0%_);(0.0%);-;"              # margin / percent rows (italic, gray, under parent)
NF_PERCENT_INPUT  = "0.0%_);(0.0%);-_);"            # assumption % inputs
NF_CAPEX_INPUT    = "$#,##0_);($#,##0)"             # capex-style $ input
NF_DATE_MONTH     = "[$-en-US]mmm-yy;@"             # month header dates
NF_CHECK          = "$#,##0.000_);($#,##0.000);-;"  # check rows (3 decimals, exposes tiny breaks)
NF_RATE_SPREAD    = '"S + "000'                      # rate-spread inputs -> renders e.g. "S + 800"

# ===========================================================================
# 7. FONTS
# ===========================================================================
FONT_BODY   = "Arial"     # every model / working tab
FONT_OUTPUT = "Segoe UI"  # output / presentation-page tabs (IS Out, UFCF Out, CIM Outputs)
FONT_SIZE   = 10.0

# ===========================================================================
# 8. PAGE SETUP (presentation tabs)
# ===========================================================================
PAGE_ORIENTATION = "portrait"
PAGE_PAPER_LETTER = "1"     # openpyxl paperSize code for Letter
MARGIN_LR = 0.7            # left / right margins (inches)
MARGIN_TB = 0.75           # top / bottom margins (inches)


# ===========================================================================
# FACTORY HELPERS
# ===========================================================================
def _argb(hex6):
    """Convert a 6-digit RGB hex (with or without '#') to opaque 8-digit ARGB."""
    h = str(hex6).lstrip("#").upper()
    return h if len(h) == 8 else "FF" + h


def font(color=SLATE_GRAY, *, bold=False, italic=False, underline=None, name=FONT_BODY, size=FONT_SIZE):
    """Build a Font. Default is the house body font: Arial 10 slate gray.
    `underline` is used only for bold+underlined section headers (e.g. the Balance
    Sheet "Assets" / "Liabilities & Equity" labels, which are NOT banner fills)."""
    return Font(name=name, size=size, color=_argb(color), bold=bold, italic=italic, underline=underline)


def fill(hex6):
    """Build a solid PatternFill from a 6-digit RGB hex."""
    a = _argb(hex6)
    return PatternFill(fill_type="solid", start_color=a, end_color=a)


NO_FILL = PatternFill(fill_type=None)


def _side(color, style="thin"):
    return Side(style=style, color=_argb(color))


CENTER = Alignment(horizontal="center")
CENTER_VCENTER = Alignment(horizontal="center", vertical="center")


def total_border():
    """The bracket that wraps a total/subtotal row: thin gray top + bottom."""
    g = _side(BORDER_GRAY)
    return Border(top=g, bottom=g)


def merge_border(existing, *, top=None, bottom=None, left=None, right=None):
    """Return a new Border that keeps `existing` sides and overlays any given ones.

    Used so that adding (say) a bottom rule to the row above a total does not wipe
    out borders that cell already carries -- and so re-running is a no-op.
    """
    return Border(
        top=top if top is not None else existing.top,
        bottom=bottom if bottom is not None else existing.bottom,
        left=left if left is not None else existing.left,
        right=right if right is not None else existing.right,
    )


# ---------------------------------------------------------------------------
# Number-format normalization (for idempotency / verification only)
# ---------------------------------------------------------------------------
# Excel may store a format like  $#,##0_);($#,##0);-;  on disk as
# "$"#,##0_);\("$"#,##0\);\-;  (quoting the $, escaping ( ) -). openpyxl writes
# our literal strings verbatim and round-trips them faithfully, but a workbook
# that was last touched by Excel can come back escaped. normalize_numfmt lets the
# verifier compare "is this effectively our format?" without false mismatches.
def normalize_numfmt(fmt):
    if fmt is None:
        return ""
    s = str(fmt)
    s = s.replace('"$"', "$").replace("\\$", "$")
    s = s.replace("\\(", "(").replace("\\)", ")")
    s = s.replace("\\-", "-").replace("\\;", ";")
    s = s.replace("[$-409]", "[$-en-US]")  # legacy LCID for en-US == [$-en-US]
    return s
